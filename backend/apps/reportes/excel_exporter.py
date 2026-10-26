import io
from decimal import ROUND_HALF_UP, Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

from apps.empresas.models import Empresa, Establecimiento, Periodo
from apps.inventario.models import Producto, Almacen
from apps.kardex.models import Incidencia
from apps.reportes.saldos_corte import resumen_saldos_al_corte
from apps.reportes.services import consultar_kardex_valorizado

FONT_NAME = "Segoe UI"

# Estilos de encabezados corporativos
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="0F172A")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, bold=False, color="475569")
HEADER_LABEL_FONT = Font(name=FONT_NAME, size=10, bold=True, color="1E293B")
HEADER_VAL_FONT = Font(name=FONT_NAME, size=10, bold=False, color="0F172A")

# Fills de grupos superiores
GROUP_OP_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
GROUP_ENT_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
GROUP_SAL_FILL = PatternFill(start_color="7C2D12", end_color="7C2D12", fill_type="solid")
GROUP_SALDO_FILL = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
GROUP_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")

# Fills de columnas
COL_HEADER_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
COL_HEADER_FONT = Font(name=FONT_NAME, size=9, bold=True, color="0F172A")

DATA_FONT = Font(name=FONT_NAME, size=9, color="0F172A")
DATA_ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
TOTALS_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
TOTALS_FONT = Font(name=FONT_NAME, size=10, bold=True, color="0F172A")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
DATA_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
TOTALS_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=Side(border_style="thin", color="0F172A"),
    bottom=Side(border_style="double", color="0F172A")
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

FMT_QTY = "#,##0.000000"
FMT_COST = "#,##0.000000"
FMT_MONEY = "#,##0.00"
FMT_DATE = "YYYY-MM-DD"



# Bloque 5: los importes van a Excel como Decimal, sin aritmética float. Excel guarda 15
# dígitos significativos (openpyxl los escribe con %.16g): se redondea de forma
# determinista (ROUND_HALF_UP) a ese límite en lugar de dejarlo a la representación binaria.
DIGITOS_EXCEL = 15


def _num(valor, cero=False):
    if valor is None:
        return Decimal('0') if cero else None
    numero = valor if isinstance(valor, Decimal) else Decimal(str(valor))
    if not numero.is_finite():
        return None
    if numero.is_zero():
        return Decimal('0')
    exceso = len(numero.as_tuple().digits) - DIGITOS_EXCEL
    if exceso > 0:
        numero = numero.quantize(Decimal(1).scaleb(numero.as_tuple().exponent + exceso), rounding=ROUND_HALF_UP)
    return numero.normalize()

def ajustar_ancho_columnas(ws, min_width=12):
    """Ajusta automáticamente el ancho de las columnas para legibilidad óptima."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignorar celdas de títulos largos combinados en filas 1 a 4 para no inflar la columna A
            if cell.row in [1, 2, 3, 4] and cell.coordinate.startswith("A"):
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_width)


def resolver_filtros_excel_kardex(*, empresa_id, periodo_id=None, producto_id=None, almacen_id=None):
    """
    SEC-003: resuelve los filtros del encabezado SOLO dentro de la empresa activa.
    Un ID de otra empresa (o inexistente) se rechaza con 400 sin revelar ningún dato.
    """
    periodo = producto = almacen = None
    errores = {}
    if periodo_id:
        periodo = Periodo.objects.filter(pk=periodo_id, empresa_id=empresa_id).first()
        if periodo is None:
            errores['periodo_id'] = 'El período no existe o no pertenece a la empresa activa.'
    if producto_id:
        producto = Producto.objects.filter(pk=producto_id, empresa_id=empresa_id).first()
        if producto is None:
            errores['producto_id'] = 'El producto no existe o no pertenece a la empresa activa.'
    if almacen_id:
        almacen = Almacen.objects.filter(pk=almacen_id, empresa_id=empresa_id).first()
        if almacen is None:
            errores['almacen_id'] = 'El almacén no existe o no pertenece a la empresa activa.'
    if errores:
        raise ValidationError(errores)
    return periodo, producto, almacen


def generar_libro_excel_kardex(
    empresa_id: int,
    periodo_id: int = None,
    producto_id: int = None,
    almacen_id: int = None,
    tipo_operacion_id: int = None,
    fecha_desde: str = None,
    fecha_hasta: str = None
) -> Workbook:
    """
    Genera el libro oficial de Excel del Kardex Valorizado con 4 hojas:
    1. Kardex Valorizado (1:1 con la vista en pantalla)
    2. Movimientos (detalle técnico completo)
    3. Incidencias (incidencias de período/empresa)
    4. Resumen de Saldos (resumen por producto y establecimiento)
    """
    periodo_filtro, producto_filtro, almacen_filtro = resolver_filtros_excel_kardex(
        empresa_id=empresa_id,
        periodo_id=periodo_id,
        producto_id=producto_id,
        almacen_id=almacen_id,
    )

    empresa = Empresa.objects.filter(pk=empresa_id).first()
    razon_social = empresa.razon_social if empresa else f"Empresa #{empresa_id}"
    ruc = empresa.ruc if empresa else "S/N"

    # Obtener la data con la misma función contable que la API
    data_kardex = consultar_kardex_valorizado(
        empresa_id=empresa_id,
        producto_id=producto_id,
        periodo_id=periodo_id,
        almacen_id=almacen_id,
        tipo_operacion_id=tipo_operacion_id,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
    )
    movimientos = data_kardex['movimientos']
    resumen = data_kardex['resumen']

    wb = Workbook()

    # =========================================================================
    # HOJA 1: Kardex Valorizado
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Kardex Valorizado"
    ws1.views.sheetView[0].showGridLines = True

    # Bloque de Encabezado Corporativo
    ws1["A1"] = f"{razon_social} — RUC: {ruc}"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = "FORMATO 13.1 — REGISTRO DE INVENTARIO PERMANENTE VALORIZADO"
    ws1["A2"].font = SUBTITLE_FONT

    # Filtros informativos
    periodo_txt = "Todos"
    if periodo_filtro:
        periodo_txt = f"{periodo_filtro.anio}-{periodo_filtro.mes:02d}"

    prod_txt = "Todos los productos"
    if producto_filtro:
        prod_txt = f"{producto_filtro.codigo_interno} - {producto_filtro.descripcion}"

    almacen_txt = "Todos los almacenes"
    if almacen_filtro:
        almacen_txt = f"{almacen_filtro.codigo} - {almacen_filtro.nombre}"

    ws1["A3"] = (
        f"Período: {periodo_txt} | Producto: {prod_txt} | Almacén: {almacen_txt}"
        + (" | Saldo valorizado: stock del almacén × CPP global" if almacen_filtro else "")
    )
    ws1["A3"].font = HEADER_VAL_FONT
    ws1["A4"] = "Método de Valuación: PROMEDIO PONDERADO (CPP GLOBAL EN PEN)"
    ws1["A4"].font = SUBTITLE_FONT

    # Fila 6: Grupos de encabezados
    ws1.merge_cells("A6:E6")
    ws1["A6"] = "DATOS DE LA OPERACIÓN"
    for col_idx in range(1, 6):
        cell = ws1.cell(row=6, column=col_idx)
        cell.fill = GROUP_OP_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER

    ws1.merge_cells("F6:H6")
    ws1["F6"] = "ENTRADAS"
    for col_idx in range(6, 9):
        cell = ws1.cell(row=6, column=col_idx)
        cell.fill = GROUP_ENT_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER

    ws1.merge_cells("I6:K6")
    ws1["I6"] = "SALIDAS"
    for col_idx in range(9, 12):
        cell = ws1.cell(row=6, column=col_idx)
        cell.fill = GROUP_SAL_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER

    ws1.merge_cells("L6:N6")
    ws1["L6"] = "SALDO VALORIZADO"
    for col_idx in range(12, 15):
        cell = ws1.cell(row=6, column=col_idx)
        cell.fill = GROUP_SALDO_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER

    # Fila 7: Subcolumnas contables
    headers_col = [
        "Fecha", "Documento Ref.", "Tipo Operación", "Naturaleza", "Almacén",
        "Cantidad", "Costo Unit.", "Valor Total",
        "Cantidad", "Costo Unit.", "Valor Total",
        "Cantidad", "CPP", "Valor Total"
    ]
    for idx, header in enumerate(headers_col, start=1):
        cell = ws1.cell(row=7, column=idx, value=header)
        cell.fill = COL_HEADER_FILL
        cell.font = COL_HEADER_FONT
        cell.alignment = ALIGN_CENTER if idx in [1, 2, 4] else (ALIGN_RIGHT if idx >= 6 else ALIGN_LEFT)
        cell.border = DATA_BORDER

    # Filas de datos
    current_row = 8

    # CON-014: con fecha inicial o período, el Kardex arranca del saldo contable previo
    # al rango (mismo valor que la API: resumen.saldo_inicial_*).
    if fecha_desde or periodo_id:
        saldos_ini = data_kardex['saldos_iniciales']
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=11)
        ws1.cell(row=current_row, column=1, value="SALDO INICIAL").alignment = Alignment(horizontal="right", vertical="center")
        cpp_ini = saldos_ini[0]['cpp'] if len(saldos_ini) == 1 else None
        for col, valor, fmt in (
            (12, resumen['saldo_inicial_cantidad'], FMT_QTY),
            (13, cpp_ini, FMT_COST),
            (14, resumen['saldo_inicial_valor'], FMT_MONEY),
        ):
            celda = ws1.cell(row=current_row, column=col, value=_num(valor))
            celda.number_format = fmt
            celda.alignment = ALIGN_RIGHT
        for col_idx in range(1, 15):
            celda = ws1.cell(row=current_row, column=col_idx)
            celda.font = TOTALS_FONT
            celda.border = DATA_BORDER
        current_row += 1

    for row_data in movimientos:
        fill_to_use = DATA_ALT_FILL if current_row % 2 == 0 else PatternFill(fill_type=None)

        # Celdas de operación
        ws1.cell(row=current_row, column=1, value=str(row_data['fecha'])).alignment = ALIGN_CENTER
        ws1.cell(row=current_row, column=2, value=row_data['documento_referencia']).alignment = ALIGN_CENTER
        ws1.cell(row=current_row, column=3, value=f"{row_data['tipo_operacion_codigo']} {row_data['tipo_operacion_descripcion']}").alignment = ALIGN_LEFT
        ws1.cell(row=current_row, column=4, value=row_data['naturaleza']).alignment = ALIGN_CENTER
        ws1.cell(row=current_row, column=5, value=f"{row_data['almacen_codigo']} - {row_data['almacen_nombre']}").alignment = ALIGN_LEFT

        # Bloque Entradas
        c_ec = ws1.cell(row=current_row, column=6, value=_num(row_data['entrada_cantidad']))
        c_ec.number_format = FMT_QTY
        c_ec.alignment = ALIGN_RIGHT

        c_ecu = ws1.cell(row=current_row, column=7, value=_num(row_data['entrada_costo_unitario']))
        c_ecu.number_format = FMT_COST
        c_ecu.alignment = ALIGN_RIGHT

        c_ev = ws1.cell(row=current_row, column=8, value=_num(row_data['entrada_valor']))
        c_ev.number_format = FMT_MONEY
        c_ev.alignment = ALIGN_RIGHT

        # Bloque Salidas
        c_sc = ws1.cell(row=current_row, column=9, value=_num(row_data['salida_cantidad']))
        c_sc.number_format = FMT_QTY
        c_sc.alignment = ALIGN_RIGHT

        c_scu = ws1.cell(row=current_row, column=10, value=_num(row_data['salida_costo_unitario']))
        c_scu.number_format = FMT_COST
        c_scu.alignment = ALIGN_RIGHT

        c_sv = ws1.cell(row=current_row, column=11, value=_num(row_data['salida_valor']))
        c_sv.number_format = FMT_MONEY
        c_sv.alignment = ALIGN_RIGHT

        # Bloque Saldo Final
        c_sdc = ws1.cell(row=current_row, column=12, value=_num(row_data['saldo_cantidad']))
        c_sdc.number_format = FMT_QTY
        c_sdc.alignment = ALIGN_RIGHT

        c_cpp = ws1.cell(row=current_row, column=13, value=_num(row_data['cpp']))
        c_cpp.number_format = FMT_COST
        c_cpp.alignment = ALIGN_RIGHT

        c_sdv = ws1.cell(row=current_row, column=14, value=_num(row_data['saldo_valor']))
        c_sdv.number_format = FMT_MONEY
        c_sdv.alignment = ALIGN_RIGHT

        for col_idx in range(1, 15):
            cell = ws1.cell(row=current_row, column=col_idx)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            if fill_to_use.fill_type:
                cell.fill = fill_to_use

        current_row += 1

    # Fila de Totales
    ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    total_lbl = ws1.cell(row=current_row, column=1, value="TOTALES DEL REPORTE")
    total_lbl.font = TOTALS_FONT
    total_lbl.alignment = Alignment(horizontal="right", vertical="center")

    tot_ec = ws1.cell(row=current_row, column=6, value=_num(resumen['total_entradas_cantidad']))
    tot_ec.number_format = FMT_QTY
    tot_ec.alignment = ALIGN_RIGHT
    tot_ec.font = TOTALS_FONT

    ws1.cell(row=current_row, column=7, value="—").alignment = ALIGN_CENTER

    tot_ev = ws1.cell(row=current_row, column=8, value=_num(resumen['total_entradas_valor']))
    tot_ev.number_format = FMT_MONEY
    tot_ev.alignment = ALIGN_RIGHT
    tot_ev.font = TOTALS_FONT

    tot_sc = ws1.cell(row=current_row, column=9, value=_num(resumen['total_salidas_cantidad']))
    tot_sc.number_format = FMT_QTY
    tot_sc.alignment = ALIGN_RIGHT
    tot_sc.font = TOTALS_FONT

    ws1.cell(row=current_row, column=10, value="—").alignment = ALIGN_CENTER

    tot_sv = ws1.cell(row=current_row, column=11, value=_num(resumen['total_salidas_valor']))
    tot_sv.number_format = FMT_MONEY
    tot_sv.alignment = ALIGN_RIGHT
    tot_sv.font = TOTALS_FONT

    tot_sdc = ws1.cell(row=current_row, column=12, value=_num(resumen['saldo_final_cantidad']))
    tot_sdc.number_format = FMT_QTY
    tot_sdc.alignment = ALIGN_RIGHT
    tot_sdc.font = TOTALS_FONT

    tot_cpp = ws1.cell(row=current_row, column=13, value=_num(resumen['cpp_final']))
    tot_cpp.number_format = FMT_COST
    tot_cpp.alignment = ALIGN_RIGHT
    tot_cpp.font = TOTALS_FONT

    tot_sdv = ws1.cell(row=current_row, column=14, value=_num(resumen['saldo_final_valor']))
    tot_sdv.number_format = FMT_MONEY
    tot_sdv.alignment = ALIGN_RIGHT
    tot_sdv.font = TOTALS_FONT

    for col_idx in range(1, 15):
        cell = ws1.cell(row=current_row, column=col_idx)
        cell.fill = TOTALS_FILL
        cell.border = TOTALS_BORDER

    ajustar_ancho_columnas(ws1)

    # =========================================================================
    # HOJA 2: Movimientos (Detalle Técnico Completo)
    # =========================================================================
    ws2 = wb.create_sheet(title="Movimientos")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = f"{razon_social} — DETALLE TÉCNICO DE MOVIMIENTOS DE KARDEX"
    ws2["A1"].font = TITLE_FONT

    mov_headers = [
        "ID Mov.", "Fecha", "Prioridad", "Secuencia", "Cód. Producto", "Descripción Producto",
        "Cód. Operación", "Operación", "Naturaleza", "Cód. Almacén", "Almacén", "Doc. Referencia",
        "Cant. Base", "Costo Unit. PEN", "Valor Mov. PEN", "Saldo Cant. Empresa",
        "CPP Empresa", "Saldo Valor Empresa", "Es Extorno", "Ajuste Motor", "Creado En"
    ]
    for idx, h in enumerate(mov_headers, start=1):
        cell = ws2.cell(row=3, column=idx, value=h)
        cell.fill = GROUP_OP_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = DATA_BORDER

    row_idx = 4
    for m in movimientos:
        ws2.cell(row=row_idx, column=1, value=m['id']).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=2, value=str(m['fecha'])).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=3, value=m['prioridad_dia']).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=4, value=m['secuencia']).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=5, value=m['producto_codigo']).alignment = ALIGN_LEFT
        ws2.cell(row=row_idx, column=6, value=m['producto_descripcion']).alignment = ALIGN_LEFT
        ws2.cell(row=row_idx, column=7, value=m['tipo_operacion_codigo']).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=8, value=m['tipo_operacion_descripcion']).alignment = ALIGN_LEFT
        ws2.cell(row=row_idx, column=9, value=m['naturaleza']).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=10, value=m['almacen_codigo']).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=11, value=m['almacen_nombre']).alignment = ALIGN_LEFT
        ws2.cell(row=row_idx, column=12, value=m['documento_referencia']).alignment = ALIGN_CENTER

        # Bloque de valores
        cant_val = m['entrada_cantidad'] if m['entrada_cantidad'] is not None else m['salida_cantidad']
        c_cb = ws2.cell(row=row_idx, column=13, value=_num(cant_val, cero=True))
        c_cb.number_format = FMT_QTY
        c_cb.alignment = ALIGN_RIGHT

        costo_val = m['entrada_costo_unitario'] if m['entrada_costo_unitario'] is not None else m['salida_costo_unitario']
        c_cu = ws2.cell(row=row_idx, column=14, value=_num(costo_val))
        c_cu.number_format = FMT_COST
        c_cu.alignment = ALIGN_RIGHT

        val_total = m['entrada_valor'] if m['entrada_valor'] is not None else m['salida_valor']
        c_vm = ws2.cell(row=row_idx, column=15, value=_num(val_total))
        c_vm.number_format = FMT_MONEY
        c_vm.alignment = ALIGN_RIGHT

        c_sc = ws2.cell(row=row_idx, column=16, value=_num(m['saldo_cantidad']))
        c_sc.number_format = FMT_QTY
        c_sc.alignment = ALIGN_RIGHT

        c_cpp = ws2.cell(row=row_idx, column=17, value=_num(m['cpp']))
        c_cpp.number_format = FMT_COST
        c_cpp.alignment = ALIGN_RIGHT

        c_sv = ws2.cell(row=row_idx, column=18, value=_num(m['saldo_valor']))
        c_sv.number_format = FMT_MONEY
        c_sv.alignment = ALIGN_RIGHT

        ws2.cell(row=row_idx, column=19, value="SÍ" if m['es_extorno'] else "NO").alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=20, value="SÍ" if m['generado_por_motor'] else "NO").alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=21, value=str(m['creado_en'])[:19] if m['creado_en'] else "").alignment = ALIGN_CENTER

        for c_idx in range(1, 22):
            cell = ws2.cell(row=row_idx, column=c_idx)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            if row_idx % 2 == 0:
                cell.fill = DATA_ALT_FILL

        row_idx += 1

    ajustar_ancho_columnas(ws2)

    # =========================================================================
    # HOJA 3: Incidencias
    # =========================================================================
    ws3 = wb.create_sheet(title="Incidencias")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = f"{razon_social} — REGISTRO DE INCIDENCIAS DE KARDEX"
    ws3["A1"].font = TITLE_FONT

    inc_headers = [
        "ID", "Período", "Documento", "Producto", "Tipo Incidencia",
        "Severidad", "Bloquea Cierre", "Estado", "Descripción", "Detectado En"
    ]
    for idx, h in enumerate(inc_headers, start=1):
        cell = ws3.cell(row=3, column=idx, value=h)
        cell.fill = GROUP_SAL_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = DATA_BORDER

    inc_qs = Incidencia.objects.filter(empresa_id=empresa_id).select_related(
        'periodo', 'producto', 'documento', 'movimiento'
    ).order_by('-detectado_en')
    if periodo_id:
        inc_qs = inc_qs.filter(periodo_id=periodo_id)

    row_inc = 4
    for inc in inc_qs:
        ws3.cell(row=row_inc, column=1, value=inc.id).alignment = ALIGN_CENTER
        p_txt = f"{inc.periodo.anio}-{inc.periodo.mes:02d}" if inc.periodo else "—"
        ws3.cell(row=row_inc, column=2, value=p_txt).alignment = ALIGN_CENTER

        doc_txt = "—"
        if inc.documento:
            doc_txt = f"{inc.documento.serie_normalizada}-{inc.documento.numero_normalizado}"
        elif inc.movimiento:
            doc_txt = f"Mov #{inc.movimiento_id}"
        ws3.cell(row=row_inc, column=3, value=doc_txt).alignment = ALIGN_CENTER

        prod_txt = f"{inc.producto.codigo_interno} - {inc.producto.descripcion}" if inc.producto else "—"
        ws3.cell(row=row_inc, column=4, value=prod_txt).alignment = ALIGN_LEFT
        ws3.cell(row=row_inc, column=5, value=inc.tipo).alignment = ALIGN_LEFT
        ws3.cell(row=row_inc, column=6, value=inc.severidad).alignment = ALIGN_CENTER
        ws3.cell(row=row_inc, column=7, value="SÍ" if inc.bloquea_cierre else "NO").alignment = ALIGN_CENTER
        ws3.cell(row=row_inc, column=8, value=inc.estado).alignment = ALIGN_CENTER
        ws3.cell(row=row_inc, column=9, value=inc.descripcion).alignment = ALIGN_LEFT
        ws3.cell(row=row_inc, column=10, value=str(inc.detectado_en)[:19] if inc.detectado_en else "").alignment = ALIGN_CENTER

        for c_idx in range(1, 11):
            cell = ws3.cell(row=row_inc, column=c_idx)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            if row_inc % 2 == 0:
                cell.fill = DATA_ALT_FILL

        row_inc += 1

    ajustar_ancho_columnas(ws3)

    # =========================================================================
    # HOJA 4: Resumen de Saldos
    # =========================================================================
    ws4 = wb.create_sheet(title="Resumen de Saldos")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A1"] = f"{razon_social} — RESUMEN DE SALDOS POR PRODUCTO Y ESTABLECIMIENTO"
    ws4["A1"].font = TITLE_FONT

    saldo_headers = [
        "Cód. Producto", "Descripción", "Establecimiento", "Almacén",
        "Stock Almacén", "CPP al Corte (PEN)", "Valor Almacén (PEN)",
        "Stock Total Empresa", "Valor Total Empresa (PEN)"
    ]
    for idx, h in enumerate(saldo_headers, start=1):
        cell = ws4.cell(row=3, column=idx, value=h)
        cell.fill = GROUP_SALDO_FILL
        cell.font = GROUP_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = DATA_BORDER

    # Bloque 4: saldos válidos a la fecha de corte del reporte (período o fecha_hasta),
    # no el stock ni el CPP actuales; con período cerrado, el cierre vigente.
    resumen_corte = resumen_saldos_al_corte(
        empresa_id=empresa_id, periodo=periodo_filtro, fecha_hasta=fecha_hasta,
        producto_id=producto_id, almacen_id=almacen_id,
    )
    ws4["A2"] = (
        f"Corte: {resumen_corte['corte'].isoformat() if resumen_corte['corte'] else 'estado actual'}"
        f" | Fuente: {resumen_corte['fuente']}"
        " | Valor por almacén = cantidad × CPP global (redondeo asignado de forma determinista)"
    )
    productos_rs = Producto.objects.in_bulk({f['producto_id'] for f in resumen_corte['filas']})
    almacenes_rs = Almacen.objects.select_related('establecimiento').in_bulk(
        {f['almacen_id'] for f in resumen_corte['filas'] if f['almacen_id']}
    )
    establecimientos_rs = Establecimiento.objects.in_bulk({f['establecimiento_id'] for f in resumen_corte['filas']})
    filas_rs = sorted(resumen_corte['filas'], key=lambda f: (
        productos_rs[f['producto_id']].codigo_interno,
        almacenes_rs[f['almacen_id']].codigo if f['almacen_id'] else '~',
    ))

    row_saldo = 4
    for fila_rs in filas_rs:
        p = productos_rs[fila_rs['producto_id']]
        alm = almacenes_rs.get(fila_rs['almacen_id'])
        estab = establecimientos_rs.get(fila_rs['establecimiento_id'])

        ws4.cell(row=row_saldo, column=1, value=p.codigo_interno).alignment = ALIGN_LEFT
        ws4.cell(row=row_saldo, column=2, value=p.descripcion).alignment = ALIGN_LEFT
        ws4.cell(row=row_saldo, column=3, value=f"{estab.codigo_sunat} - {estab.nombre}" if estab else "").alignment = ALIGN_LEFT
        ws4.cell(row=row_saldo, column=4, value=(
            f"{alm.codigo} - {alm.nombre}" if alm else "(ajuste de redondeo, sin almacén)"
        )).alignment = ALIGN_LEFT

        for col, valor, fmt in (
            (5, fila_rs['cantidad_almacen'], FMT_QTY),
            (6, fila_rs['cpp'], FMT_COST),
            (7, fila_rs['valor_almacen'], FMT_MONEY),
            (8, fila_rs['cantidad_empresa'], FMT_QTY),
            (9, fila_rs['valor_empresa'], FMT_MONEY),
        ):
            celda = ws4.cell(row=row_saldo, column=col, value=_num(valor))
            celda.number_format = fmt
            celda.alignment = ALIGN_RIGHT

        for c_idx in range(1, 10):
            cell = ws4.cell(row=row_saldo, column=c_idx)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            if row_saldo % 2 == 0:
                cell.fill = DATA_ALT_FILL

        row_saldo += 1

    ajustar_ancho_columnas(ws4)

    return wb


def generar_bytes_excel_kardex(**kwargs) -> io.BytesIO:
    """Helper que invoca generar_libro_excel_kardex y entrega un stream de bytes en memoria."""
    wb = generar_libro_excel_kardex(**kwargs)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
