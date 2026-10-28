import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.reportes.excel_seguro import neutralizar_formulas
from apps.reportes.formato131_service import generar_formato_131


AZUL = '17365D'
AZUL_CLARO = 'D9EAF7'
VERDE = 'E2F0D9'
ROJO = 'FCE4D6'
BORDER = Border(*(Side(style='thin', color='A6A6A6'),) * 4)
FORMATO_NUMERO = '#,##0.00000000'


def nombre_archivo_formato_131(periodo, ruc):
    """Bloque 13: `Formato_13_1_2026-01_20123456789.xlsx` (solo dígitos, letras, '-', '_' y '.')."""
    ruc_seguro = ''.join(c for c in str(ruc or '') if c.isalnum())[:20]
    partes = ['Formato_13_1', f'{periodo.anio}-{periodo.mes:02d}'] + ([ruc_seguro] if ruc_seguro else [])
    return '_'.join(partes) + '.xlsx'


def _pie_conciliacion(hoja, fila, reporte, fuente_cierre):
    """
    Bloque 13: lo mismo que la pantalla muestra debajo de cada tabla — de dónde sale el saldo
    final (cierre vigente con versión y hash, o calculado) y la microdiferencia de cierre. Solo
    presentación: los valores vienen de generar_formato_131 sin recalcular nada.
    """
    if reporte.get('fuente_saldo_final') == 'CIERRE' and fuente_cierre:
        fuente = f"CIERRE v{fuente_cierre['version']} (hash {fuente_cierre['hash'][:12]}…)"
    else:
        fuente = 'CALCULADO (sin cierre vigente)'
    hoja.cell(fila, 1, 'FUENTE DEL SALDO FINAL').font = Font(bold=True)
    hoja.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=14)
    hoja.cell(fila, 2, fuente)
    micro = Decimal(reporte.get('microdiferencia_cierre') or '0')
    if micro == 0:
        return fila
    fila += 1
    etiqueta = (
        'AJUSTE DE REDONDEO DEL REPARTO POR ESTABLECIMIENTO (no es una diferencia contable)'
        if reporte.get('fuente_saldo_final') == 'CALCULADO'
        else 'MICRODIFERENCIA OFICIAL ASIGNADA AL ESTABLECIMIENTO PRINCIPAL (absorbida por redondeo contable)'
    )
    hoja.cell(fila, 1, etiqueta).font = Font(bold=True, color='9C5700')
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=13)
    celda = hoja.cell(fila, 14, micro)
    celda.number_format = FORMATO_NUMERO
    celda.font = Font(bold=True, color='9C5700')
    return fila


def generar_excel_formato_131(**filtros):
    data = generar_formato_131(**filtros)
    libro = Workbook()
    libro.remove(libro.active)
    for indice, reporte in enumerate(data['reportes'], start=1):
        cab = reporte['cabecera']
        hoja = libro.create_sheet(f'13.1 {indice:02d}')
        hoja.sheet_view.showGridLines = False
        hoja.page_setup.orientation = 'landscape'
        hoja.page_setup.fitToWidth = 1
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.freeze_panes = 'A13'
        hoja.merge_cells('A1:M1')
        hoja['A1'] = 'FORMATO 13.1: REGISTRO DE INVENTARIO PERMANENTE VALORIZADO'
        hoja['A1'].font = Font(bold=True, size=14, color='FFFFFF')
        hoja['A1'].fill = PatternFill('solid', fgColor=AZUL)
        hoja['A1'].alignment = Alignment(horizontal='center')
        datos = [
            ('PERÍODO', data['periodo']), ('RUC', data['empresa']['ruc']),
            ('RAZÓN SOCIAL', data['empresa']['razon_social']),
            ('ESTABLECIMIENTO', f"{cab['establecimiento']['codigo_sunat']} - {cab['establecimiento']['nombre']}"),
            ('CÓDIGO DE EXISTENCIA', cab['codigo_existencia']),
            ('TIPO DE EXISTENCIA — TABLA 5', f"{cab['tipo_existencia']['codigo']} - {cab['tipo_existencia']['descripcion']}"),
            ('DESCRIPCIÓN', cab['descripcion_existencia']),
            ('UNIDAD DE MEDIDA — TABLA 6', f"{cab['unidad_medida']['codigo']} - {cab['unidad_medida']['descripcion']}"),
            ('MÉTODO DE VALUACIÓN', cab['metodo_valuacion']),
        ]
        for fila, (etiqueta, valor) in enumerate(datos, start=2):
            hoja.cell(fila, 1, etiqueta).font = Font(bold=True)
            hoja.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=13)
            hoja.cell(fila, 2, valor)

        grupos = [('A11:A11', 'OPERACIÓN'), ('B11:D11', 'DOCUMENTO DE TRASLADO / COMPROBANTE'),
                  ('E11:E11', 'TIPO DE OPERACIÓN'), ('F11:H11', 'ENTRADAS'),
                  ('I11:K11', 'SALIDAS'), ('L11:N11', 'SALDO FINAL')]
        for rango, titulo in grupos:
            hoja.merge_cells(rango)
            celda = hoja[rango.split(':')[0]]
            celda.value = titulo
            celda.fill = PatternFill('solid', fgColor=AZUL)
            celda.font = Font(bold=True, color='FFFFFF')
            celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        encabezados = ['Fecha', 'Tipo', 'Serie', 'Número', 'Código Tabla 12',
                       'Cantidad', 'Costo Unitario', 'Costo Total',
                       'Cantidad', 'Costo Unitario', 'Costo Total',
                       'Cantidad', 'Costo Unitario', 'Costo Total']
        for columna, titulo in enumerate(encabezados, start=1):
            celda = hoja.cell(12, columna, titulo)
            celda.font = Font(bold=True)
            celda.fill = PatternFill('solid', fgColor=AZUL_CLARO)
            celda.alignment = Alignment(horizontal='center', wrap_text=True)
            celda.border = BORDER
        # CON-010: saldo inicial valorizado (cantidad, costo unitario y costo total).
        inicial = reporte.get('saldo_inicial_valorizado') or {}
        hoja.cell(13, 1, 'SALDO INICIAL').font = Font(bold=True)
        hoja.merge_cells(start_row=13, start_column=1, end_row=13, end_column=11)
        for desplazamiento, clave in enumerate(('cantidad', 'costo_unitario', 'costo_total')):
            celda = hoja.cell(13, 12 + desplazamiento, Decimal(inicial[clave]) if clave in inicial else None)
            celda.number_format = FORMATO_NUMERO
            celda.font = Font(bold=True)
            celda.fill = PatternFill('solid', fgColor=AZUL_CLARO)
        for celda in hoja[13]:
            celda.border = BORDER
        for fila, movimiento in enumerate(reporte['movimientos'], start=14):
            valores = [movimiento['fecha_emision'], movimiento['documento']['tipo_documento'],
                       movimiento['documento']['serie'], movimiento['documento']['numero'],
                       movimiento['tipo_operacion']]
            for columna, valor in enumerate(valores, start=1):
                hoja.cell(fila, columna, valor)
            for inicio, bloque, color in ((6, movimiento['entrada'], VERDE), (9, movimiento['salida'], ROJO), (12, movimiento['saldo_final'], AZUL_CLARO)):
                for desplazamiento, clave in enumerate(('cantidad', 'costo_unitario', 'costo_total')):
                    celda = hoja.cell(fila, inicio + desplazamiento, Decimal(bloque[clave]) if bloque else None)
                    celda.number_format = FORMATO_NUMERO
                    celda.fill = PatternFill('solid', fgColor=color)
            for celda in hoja[fila]:
                celda.border = BORDER
                celda.alignment = Alignment(horizontal='right' if celda.column >= 6 else 'center')
        for columna, ancho in enumerate((12, 9, 12, 14, 18, 14, 16, 16, 14, 16, 16, 14, 16, 16), start=1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        ultima = _pie_conciliacion(hoja, 15 + len(reporte['movimientos']), reporte, data.get('fuente_cierre'))
        hoja.print_title_rows = '11:12'
        hoja.print_area = f'A1:N{ultima}'
    if not data['reportes']:
        hoja = libro.create_sheet('Formato 13.1')
        hoja['A1'] = 'No existen movimientos para los filtros seleccionados.'
    neutralizar_formulas(libro)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer
