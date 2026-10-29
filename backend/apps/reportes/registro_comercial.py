"""
Bloque 16 — Registro de ventas / Registro de compras (reporte comercial interno).

NO es el PLE/SIRE oficial. Lee `documento` (sin tablas propias) del período: fecha, tercero,
tipo, serie, número, moneda, base, IGV, total, condición y estado. Las notas de crédito restan
(importes con signo negativo); los comprobantes ANULADOS se listan pero no suman.
"""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

from apps.documentos.finanzas import q8, texto
from apps.documentos.models import Documento
from apps.empresas.models import Periodo
from apps.reportes.excel_seguro import neutralizar_formulas

CERO = Decimal('0')
ESTADOS = ('CONFIRMADO', 'ANULADO', 'BORRADOR')


def generar_registro(*, empresa_id, clase, periodo_id, estado=None, tercero_id=None):
    periodo = Periodo.objects.select_related('empresa').filter(pk=periodo_id, empresa_id=empresa_id).first()
    if periodo is None:
        raise ValidationError({'periodo_id': 'El período no existe o no pertenece a la empresa activa.'})
    qs = (Documento.objects.filter(empresa_id=empresa_id, clase=clase, periodo=periodo)
          .select_related('tercero', 'tipo_comprobante', 'moneda_codigo')
          .order_by('fecha_emision', 'tipo_comprobante__codigo_sunat', 'serie_normalizada', 'id'))
    if estado:
        if estado not in ESTADOS:
            raise ValidationError({'estado': f"Use {', '.join(ESTADOS)}."})
        qs = qs.filter(estado=estado)
    else:
        qs = qs.exclude(estado='BORRADOR')
    if tercero_id:
        qs = qs.filter(tercero_id=tercero_id)
    filas, totales = [], {}
    for d in qs:
        signo = Decimal('-1') if d.tipo_comprobante.efecto_nota == 'CREDITO' else Decimal('1')
        base, igv, total = (q8((v or CERO) * signo) for v in (d.subtotal, d.igv, d.total))
        filas.append({
            'id': d.id, 'fecha_emision': d.fecha_emision, 'tipo_comprobante_codigo': d.tipo_comprobante.codigo_sunat,
            'tipo_comprobante_descripcion': d.tipo_comprobante.descripcion, 'serie': d.serie_normalizada,
            'numero': d.numero or d.numero_normalizado,
            'tercero_tipo_documento': d.tercero.tipo_documento if d.tercero else None,
            'tercero_numero_documento': d.tercero.numero_documento if d.tercero else None,
            'tercero_razon_social': d.tercero.razon_social if d.tercero else ('CLIENTES VARIOS' if clase == 'VENTA' else None),
            'moneda_codigo': d.moneda_codigo_id, 'monto_tc': str(d.monto_tc) if d.monto_tc is not None else None,
            'base': texto(base), 'igv': texto(igv), 'total': texto(total),
            'condicion_pago': d.condicion_pago, 'estado': d.estado,
        })
        if d.estado == 'CONFIRMADO':
            t = totales.setdefault(d.moneda_codigo_id, {'base': CERO, 'igv': CERO, 'total': CERO, 'comprobantes': 0})
            t['base'] += base
            t['igv'] += igv
            t['total'] += total
            t['comprobantes'] += 1
    return {
        'clase': clase,
        'periodo': {'id': periodo.id, 'anio': periodo.anio, 'mes': periodo.mes},
        'empresa': {'ruc': periodo.empresa.ruc, 'razon_social': periodo.empresa.razon_social},
        'filas': filas,
        'totales': {m: {**{k: texto(v) for k, v in t.items() if k != 'comprobantes'}, 'comprobantes': t['comprobantes']}
                    for m, t in totales.items()},
    }


def nombre_archivo_registro(datos):
    p = datos['periodo']
    tipo = 'Registro_Ventas' if datos['clase'] == 'VENTA' else 'Registro_Compras'
    ruc = ''.join(c for c in datos['empresa']['ruc'] if c.isalnum())
    return f"{tipo}_{p['anio']}-{p['mes']:02d}_{ruc}.xlsx"


def excel_registro(datos):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Registro de ventas' if datos['clase'] == 'VENTA' else 'Registro de compras'
    titulo = 'REGISTRO DE VENTAS' if datos['clase'] == 'VENTA' else 'REGISTRO DE COMPRAS'
    p = datos['periodo']
    hoja['A1'] = f"{titulo} — {p['anio']}-{p['mes']:02d} (reporte comercial interno, no es PLE/SIRE)"
    hoja['A1'].font = Font(bold=True, size=13)
    hoja['A2'] = f"{datos['empresa']['ruc']} — {datos['empresa']['razon_social']}"
    rol = 'Cliente' if datos['clase'] == 'VENTA' else 'Proveedor'
    encabezados = ['Fecha', 'Tipo', 'Serie', 'Número', f'Doc. {rol.lower()}', rol, 'Moneda', 'T.C.',
                   'Base', 'IGV', 'Total', 'Condición', 'Estado']
    for c, h in enumerate(encabezados, start=1):
        celda = hoja.cell(4, c, h)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='17365D')
        celda.alignment = Alignment(horizontal='center')
    fila = 5
    for f in datos['filas']:
        valores = [f['fecha_emision'], f['tipo_comprobante_codigo'], f['serie'], f['numero'],
                   f['tercero_numero_documento'], f['tercero_razon_social'], f['moneda_codigo'],
                   Decimal(f['monto_tc']) if f['monto_tc'] else None,
                   Decimal(f['base']), Decimal(f['igv']), Decimal(f['total']), f['condicion_pago'], f['estado']]
        for c, v in enumerate(valores, start=1):
            celda = hoja.cell(fila, c, v)
            if c in (8, 9, 10, 11):
                celda.number_format = '#,##0.00000000' if c == 8 else '#,##0.00'
        fila += 1
    fila += 1
    for moneda, t in sorted(datos['totales'].items()):
        hoja.cell(fila, 6, f'TOTAL {moneda} ({t["comprobantes"]} confirmados)').font = Font(bold=True)
        for c, k in ((9, 'base'), (10, 'igv'), (11, 'total')):
            celda = hoja.cell(fila, c, Decimal(t[k]))
            celda.number_format = '#,##0.00'
            celda.font = Font(bold=True)
        fila += 1
    for c, ancho in enumerate((12, 6, 8, 12, 14, 40, 8, 12, 14, 14, 14, 11, 12), start=1):
        hoja.column_dimensions[get_column_letter(c)].width = ancho
    hoja.freeze_panes = 'A5'
    neutralizar_formulas(libro)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer
