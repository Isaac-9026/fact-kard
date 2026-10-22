import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient

from apps.catalogos.models import Moneda, TipoComprobante, TipoExistencia, TipoOperacion, UnidadMedida
from apps.documentos.models import Documento, DocumentoDetalle
from apps.empresas.models import ConfigEmpresa, Empresa, Establecimiento, Periodo
from apps.inventario.models import Almacen, Producto, ProductoUnidad
from apps.kardex.models import Movimiento, ValorizacionMovimiento
from apps.usuarios.models import Usuario, UsuarioEmpresa


class Formato131Test(TestCase):
    @classmethod
    def setUpTestData(cls):
        moneda = Moneda.objects.create(codigo='PEN', nombre='Soles')
        cls.empresa = Empresa.objects.create(ruc='20111111111', razon_social='EMPRESA REPORTE', moneda_base_codigo=moneda)
        cls.otra = Empresa.objects.create(ruc='20222222222', razon_social='OTRA EMPRESA', moneda_base_codigo=moneda)
        cls.est_a = Establecimiento.objects.create(empresa=cls.empresa, codigo_sunat='0000', nombre='PRINCIPAL')
        cls.est_b = Establecimiento.objects.create(empresa=cls.empresa, codigo_sunat='0001', nombre='SUCURSAL')
        cls.est_otra = Establecimiento.objects.create(empresa=cls.otra, codigo_sunat='0000', nombre='AJENA')
        ConfigEmpresa.objects.create(empresa=cls.empresa, establecimiento_principal=cls.est_a, metodo_valuacion='CPP')
        cls.alm_a = Almacen.objects.create(empresa=cls.empresa, establecimiento=cls.est_a, codigo='A', nombre='ALM A')
        cls.alm_b = Almacen.objects.create(empresa=cls.empresa, establecimiento=cls.est_b, codigo='B', nombre='ALM B')
        tipo = TipoExistencia.objects.create(codigo='01', descripcion='MERCADERÍAS')
        unidad = UnidadMedida.objects.create(codigo='NIU', descripcion='UNIDADES')
        cls.producto = Producto.objects.create(empresa=cls.empresa, codigo_interno='P-001', descripcion='PRODUCTO UNO', tipo_existencia_codigo=tipo)
        cls.producto_otro = Producto.objects.create(empresa=cls.otra, codigo_interno='P-X', descripcion='AJENO', tipo_existencia_codigo=tipo)
        ProductoUnidad.objects.create(producto=cls.producto, unidad_codigo=unidad, factor_a_base=Decimal('1'), es_base=True)
        cls.periodo = Periodo.objects.create(empresa=cls.empresa, anio=2026, mes=10, orden=202610, fecha_inicio=date(2026, 10, 1), fecha_fin=date(2026, 10, 31))
        cls.periodo_otro = Periodo.objects.create(empresa=cls.otra, anio=2026, mes=10, orden=202610, fecha_inicio=date(2026, 10, 1), fecha_fin=date(2026, 10, 31))
        cls.op_entrada = TipoOperacion.objects.create(codigo_interno='COMPRA_F131', codigo_sunat='02', descripcion='COMPRA', naturaleza='ENTRADA', origen_requerido='DOCUMENTO')
        cls.op_salida = TipoOperacion.objects.create(codigo_interno='TR_SAL_F131', codigo_sunat='11', descripcion='TRANSFERENCIA SALIDA', naturaleza='SALIDA', origen_requerido='TRANSFERENCIA')
        cls.op_tr_entrada = TipoOperacion.objects.create(codigo_interno='TR_ENT_F131', codigo_sunat='21', descripcion='TRANSFERENCIA ENTRADA', naturaleza='ENTRADA', origen_requerido='TRANSFERENCIA')
        comprobante = TipoComprobante.objects.create(codigo_sunat='01', descripcion='FACTURA')
        doc = Documento.objects.create(empresa=cls.empresa, clase='COMPRA', periodo=cls.periodo, ruc_emisor='20999999999', tipo_comprobante=comprobante, serie='F001', serie_normalizada='F001', numero='10', numero_normalizado='10', fecha_emision=date(2026, 10, 21), moneda_codigo=moneda)
        detalle = DocumentoDetalle.objects.create(documento=doc, linea=1, producto=cls.producto, almacen=cls.alm_a, tipo_operacion=cls.op_entrada, cantidad_original=Decimal('100'), factor_conversion=Decimal('1'), cantidad_base=Decimal('100'))
        cls.compra = cls._mov(cls.alm_a, cls.op_entrada, 'ENTRADA', Decimal('100'), 2, 1, detalle=detalle, saldo=Decimal('100'))
        cls.tr_ent = cls._mov(cls.alm_b, cls.op_tr_entrada, 'ENTRADA', Decimal('10'), 2, 2, saldo=Decimal('100'))
        cls.tr_sal = cls._mov(cls.alm_a, cls.op_salida, 'SALIDA', Decimal('10'), 4, 1, saldo=Decimal('100'))
        cls.user = Usuario.objects.create(username='reporte', email='r@test.local', activo=True)
        UsuarioEmpresa.objects.create(usuario=cls.user, empresa=cls.empresa, rol='USUARIO', activo=True)

    @classmethod
    def _mov(cls, almacen, operacion, naturaleza, cantidad, prioridad, secuencia, detalle=None, saldo=Decimal('0')):
        mov = Movimiento.objects.create(empresa=cls.empresa, producto=cls.producto, almacen=almacen, periodo=cls.periodo, documento_detalle=detalle, tipo_operacion=operacion, fecha_movimiento=date(2026, 10, 21), prioridad_dia=prioridad, secuencia=secuencia, naturaleza=naturaleza, cantidad_base=cantidad)
        ValorizacionMovimiento.objects.create(movimiento=mov, costo_unitario_pen=Decimal('10'), valor_movimiento_pen=cantidad * Decimal('10'), saldo_cantidad_empresa=saldo, saldo_valor_empresa=saldo * Decimal('10'), cpp_empresa=Decimal('10'))
        return mov

    def setUp(self):
        self.client = APIClient(); self.client.force_authenticate(self.user)

    def get(self, params=None, empresa=None):
        return self.client.get(reverse('formato-131'), params or {'periodo_id': self.periodo.id}, HTTP_X_EMPRESA_ID=str((empresa or self.empresa).id))

    def test_01_requiere_empresa_activa(self):
        self.assertEqual(self.client.get(reverse('formato-131'), {'periodo_id': self.periodo.id}).status_code, 403)

    def test_02_empresa_ajena_es_rechazada(self):
        self.assertEqual(self.get(empresa=self.otra).status_code, 403)

    def test_03_periodo_debe_ser_de_empresa(self):
        self.assertEqual(self.get({'periodo_id': self.periodo_otro.id}).status_code, 400)

    def test_04_establecimiento_debe_ser_de_empresa(self):
        self.assertEqual(self.get({'periodo_id': self.periodo.id, 'establecimiento_id': self.est_otra.id}).status_code, 400)

    def test_05_producto_debe_ser_de_empresa(self):
        self.assertEqual(self.get({'periodo_id': self.periodo.id, 'producto_id': self.producto_otro.id}).status_code, 400)

    def test_06_cabecera_ruc(self): self.assertEqual(self.get().data['empresa']['ruc'], self.empresa.ruc)
    def test_07_cabecera_razon_social(self): self.assertEqual(self.get().data['empresa']['razon_social'], self.empresa.razon_social)
    def test_08_codigo_establecimiento(self): self.assertEqual(self.get({'periodo_id': self.periodo.id, 'establecimiento_id': self.est_a.id}).data['reportes'][0]['cabecera']['establecimiento']['codigo_sunat'], '0000')
    def test_09_tabla_5(self): self.assertEqual(self.get().data['reportes'][0]['cabecera']['tipo_existencia']['codigo'], '01')
    def test_10_tabla_6(self): self.assertEqual(self.get().data['reportes'][0]['cabecera']['unidad_medida']['codigo'], 'NIU')
    def test_11_tabla_10_documento(self): self.assertEqual(self.get({'periodo_id': self.periodo.id, 'establecimiento_id': self.est_a.id}).data['reportes'][0]['movimientos'][0]['documento']['tipo_documento'], '01')
    def test_12_operacion_interna_documento_00(self): self.assertTrue(any(m['documento']['tipo_documento'] == '00' for r in self.get().data['reportes'] for m in r['movimientos']))
    def test_13_tabla_12(self): self.assertEqual(self.get({'periodo_id': self.periodo.id, 'establecimiento_id': self.est_a.id}).data['reportes'][0]['movimientos'][0]['tipo_operacion'], '02')
    def test_14_orden_oficial(self): self.assertEqual([m['movimiento_id'] for r in self.get().data['reportes'] for m in r['movimientos'] if r['cabecera']['establecimiento']['id'] == self.est_a.id], [self.compra.id, self.tr_sal.id])
    def test_15_cpp_es_global(self): self.assertTrue(all(m['saldo_final']['costo_unitario'] == '10.00000000' for r in self.get().data['reportes'] for m in r['movimientos']))
    def test_16_saldo_fisico_establecimiento(self): self.assertEqual(self.get({'periodo_id': self.periodo.id, 'establecimiento_id': self.est_a.id}).data['reportes'][0]['movimientos'][-1]['saldo_final']['cantidad'], '90.00000000')
    def test_17_transferencia_redistribuye(self): self.assertEqual(self.get({'periodo_id': self.periodo.id, 'establecimiento_id': self.est_b.id}).data['reportes'][0]['movimientos'][-1]['saldo_final']['cantidad'], '10.00000000')
    def test_18_transferencia_es_neutral_global(self):
        finales = [Decimal(r['movimientos'][-1]['saldo_final']['cantidad']) for r in self.get().data['reportes']]
        self.assertEqual(sum(finales), Decimal('100'))

    def excel(self): return self.client.get(reverse('formato-131-excel'), {'periodo_id': self.periodo.id}, HTTP_X_EMPRESA_ID=str(self.empresa.id))
    def test_19_excel_200(self): self.assertEqual(self.excel().status_code, 200)
    def test_20_excel_content_type(self): self.assertIn('spreadsheetml', self.excel()['Content-Type'])
    def test_21_excel_abre(self): self.assertGreater(len(openpyxl.load_workbook(io.BytesIO(self.excel().content)).sheetnames), 0)
    def test_22_excel_encabezado_oficial(self): self.assertIn('FORMATO 13.1', openpyxl.load_workbook(io.BytesIO(self.excel().content)).active['A1'].value)
    def test_23_excel_bloques_combinados(self): self.assertIn('B11:D11', {str(r) for r in openpyxl.load_workbook(io.BytesIO(self.excel().content)).active.merged_cells.ranges})
    def test_24_excel_multiempresa(self):
        respuesta = self.client.get(reverse('formato-131-excel'), {'periodo_id': self.periodo_otro.id}, HTTP_X_EMPRESA_ID=str(self.empresa.id))
        self.assertEqual(respuesta.status_code, 400)
