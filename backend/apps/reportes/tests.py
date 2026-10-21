import io
from datetime import date
from decimal import Decimal
import openpyxl
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient

from apps.catalogos.models import Moneda, TipoExistencia, UnidadMedida, TipoOperacion, TipoComprobante
from apps.empresas.models import Empresa, Establecimiento, Periodo
from apps.inventario.models import Almacen, Producto
from apps.documentos.models import Documento, DocumentoDetalle
from apps.kardex.models import Movimiento, ValorizacionMovimiento, SaldoCppActual, Incidencia, StockAlmacen
from apps.usuarios.models import Usuario, UsuarioEmpresa


class KardexReportesTest(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Moneda
        self.moneda = Moneda.objects.create(codigo='PEN', nombre='Soles')

        # Empresa A
        self.empresa_a = Empresa.objects.create(
            ruc='20100070970',
            razon_social='EMPRESA A S.A.C.',
            moneda_base_codigo=self.moneda
        )
        self.estab_a = Establecimiento.objects.create(
            empresa=self.empresa_a,
            codigo_sunat='0000',
            nombre='SEDE CENTRAL A'
        )
        self.almacen_a1 = Almacen.objects.create(
            empresa=self.empresa_a,
            establecimiento=self.estab_a,
            codigo='ALM01',
            nombre='ALMACÉN PRINCIPAL A'
        )
        self.periodo_a = Periodo.objects.create(
            empresa=self.empresa_a,
            anio=2026,
            mes=1,
            orden=202601,
            fecha_inicio=date(2026, 1, 1),
            fecha_fin=date(2026, 1, 31),
            estado='ABIERTO'
        )

        # Empresa B
        self.empresa_b = Empresa.objects.create(
            ruc='20200080980',
            razon_social='EMPRESA B S.A.C.',
            moneda_base_codigo=self.moneda
        )
        self.estab_b = Establecimiento.objects.create(
            empresa=self.empresa_b,
            codigo_sunat='0000',
            nombre='SEDE CENTRAL B'
        )
        self.almacen_b = Almacen.objects.create(
            empresa=self.empresa_b,
            establecimiento=self.estab_b,
            codigo='ALM-B',
            nombre='ALMACÉN B'
        )
        self.periodo_b = Periodo.objects.create(
            empresa=self.empresa_b,
            anio=2026,
            mes=1,
            orden=202601,
            fecha_inicio=date(2026, 1, 1),
            fecha_fin=date(2026, 1, 31),
            estado='ABIERTO'
        )

        # Catálogos base
        self.tipo_existencia = TipoExistencia.objects.create(codigo='01', descripcion='MERCADERIAS')
        self.op_compra = TipoOperacion.objects.create(
            codigo_interno='COMPRA_NACIONAL',
            codigo_sunat='02',
            descripcion='COMPRA NACIONAL',
            naturaleza='ENTRADA',
            origen_requerido='DOCUMENTO',
        )
        self.op_venta = TipoOperacion.objects.create(
            codigo_interno='VENTA_NACIONAL',
            codigo_sunat='01',
            descripcion='VENTA NACIONAL',
            naturaleza='SALIDA',
            origen_requerido='DOCUMENTO',
        )

        # Productos
        self.prod_a = Producto.objects.create(
            empresa=self.empresa_a,
            tipo_existencia_codigo=self.tipo_existencia,
            codigo_interno='PROD-A',
            descripcion='PRODUCTO A TEST'
        )
        self.prod_b = Producto.objects.create(
            empresa=self.empresa_b,
            tipo_existencia_codigo=self.tipo_existencia,
            codigo_interno='PROD-B',
            descripcion='PRODUCTO B TEST'
        )

        # Usuarios
        self.user_a = Usuario.objects.create(
            username='user_a',
            email='usera@factkard.pe',
            password_hash='hash',
            activo=True
        )
        UsuarioEmpresa.objects.create(
            usuario=self.user_a,
            empresa=self.empresa_a,
            rol='ADMIN',
            activo=True
        )

        self.user_b = Usuario.objects.create(
            username='user_b',
            email='userb@factkard.pe',
            password_hash='hash',
            activo=True
        )
        UsuarioEmpresa.objects.create(
            usuario=self.user_b,
            empresa=self.empresa_b,
            rol='ADMIN',
            activo=True
        )

        # Documento origen para Mov 1
        self.tipo_factura = TipoComprobante.objects.create(
            codigo_sunat='01',
            descripcion='FACTURA',
            activo=True
        )
        self.doc_compra = Documento.objects.create(
            empresa=self.empresa_a,
            clase='COMPRA',
            periodo=self.periodo_a,
            ruc_emisor='20500000001',
            tipo_comprobante=self.tipo_factura,
            serie='F001',
            serie_normalizada='F001',
            numero='1234',
            numero_normalizado='00001234',
            fecha_emision=date(2026, 1, 10),
            moneda_codigo=self.moneda,
        )
        self.doc_detalle_compra = DocumentoDetalle.objects.create(
            documento=self.doc_compra,
            linea=1,
            producto=self.prod_a,
            almacen=self.almacen_a1,
            tipo_operacion=self.op_compra,
        )

        # Movimientos cronológicos en Empresa A
        # Mov 1: Compra inicial de 100 unidades a S/ 10.00
        self.mov1 = Movimiento.objects.create(
            empresa=self.empresa_a,
            producto=self.prod_a,
            almacen=self.almacen_a1,
            periodo=self.periodo_a,
            tipo_operacion=self.op_compra,
            documento_detalle=self.doc_detalle_compra,
            fecha_movimiento=date(2026, 1, 10),
            prioridad_dia=2,
            secuencia=1,
            naturaleza='ENTRADA',
            cantidad_base=Decimal('100.00000000'),
        )
        ValorizacionMovimiento.objects.create(
            movimiento=self.mov1,
            costo_unitario_pen=Decimal('10.00000000'),
            valor_movimiento_pen=Decimal('1000.00000000'),
            saldo_cantidad_empresa=Decimal('100.00000000'),
            saldo_valor_empresa=Decimal('1000.00000000'),
            cpp_empresa=Decimal('10.00000000'),
        )

        # Mov 2: Segunda compra de 50 unidades a S/ 13.00 (mismo día posterior, o siguiente día)
        self.mov2 = Movimiento.objects.create(
            empresa=self.empresa_a,
            producto=self.prod_a,
            almacen=self.almacen_a1,
            periodo=self.periodo_a,
            tipo_operacion=self.op_compra,
            fecha_movimiento=date(2026, 1, 15),
            prioridad_dia=2,
            secuencia=1,
            naturaleza='ENTRADA',
            cantidad_base=Decimal('50.00000000'),
        )
        # Saldo: 150 und, Valor: 1000 + 650 = 1650, CPP = 1650/150 = 11.00
        ValorizacionMovimiento.objects.create(
            movimiento=self.mov2,
            costo_unitario_pen=Decimal('13.00000000'),
            valor_movimiento_pen=Decimal('650.00000000'),
            saldo_cantidad_empresa=Decimal('150.00000000'),
            saldo_valor_empresa=Decimal('1650.00000000'),
            cpp_empresa=Decimal('11.00000000'),
        )

        # Mov 3: Venta de 40 unidades a CPP S/ 11.00
        self.mov3 = Movimiento.objects.create(
            empresa=self.empresa_a,
            producto=self.prod_a,
            almacen=self.almacen_a1,
            periodo=self.periodo_a,
            tipo_operacion=self.op_venta,
            fecha_movimiento=date(2026, 1, 20),
            prioridad_dia=4,
            secuencia=1,
            naturaleza='SALIDA',
            cantidad_base=Decimal('40.00000000'),
        )
        # Saldo: 110 und, Salida valor: 440, Saldo valor: 1210, CPP: 11.00
        ValorizacionMovimiento.objects.create(
            movimiento=self.mov3,
            costo_unitario_pen=Decimal('11.00000000'),
            valor_movimiento_pen=Decimal('440.00000000'),
            saldo_cantidad_empresa=Decimal('110.00000000'),
            saldo_valor_empresa=Decimal('1210.00000000'),
            cpp_empresa=Decimal('11.00000000'),
        )

        # Stock y CPP actual en Empresa A
        StockAlmacen.objects.create(
            empresa=self.empresa_a,
            almacen=self.almacen_a1,
            producto=self.prod_a,
            cantidad_actual=Decimal('110.00000000')
        )
        SaldoCppActual.objects.create(
            empresa=self.empresa_a,
            producto=self.prod_a,
            cantidad_total=Decimal('110.00000000'),
            valor_total_pen=Decimal('1210.00000000'),
            cpp_pen=Decimal('11.00000000'),
            ultimo_movimiento=self.mov3
        )

        # Movimiento en Empresa B
        self.mov_b = Movimiento.objects.create(
            empresa=self.empresa_b,
            producto=self.prod_b,
            almacen=self.almacen_b,
            periodo=self.periodo_b,
            tipo_operacion=self.op_compra,
            fecha_movimiento=date(2026, 1, 5),
            prioridad_dia=2,
            secuencia=1,
            naturaleza='ENTRADA',
            cantidad_base=Decimal('200.00000000'),
        )
        ValorizacionMovimiento.objects.create(
            movimiento=self.mov_b,
            costo_unitario_pen=Decimal('20.00000000'),
            valor_movimiento_pen=Decimal('4000.00000000'),
            saldo_cantidad_empresa=Decimal('200.00000000'),
            saldo_valor_empresa=Decimal('4000.00000000'),
            cpp_empresa=Decimal('20.00000000'),
        )

    def test_kardex_reporte_orden_cronologico_y_estructura(self):
        """Verifica que el endpoint devuelva las filas en el orden cronológico exacto y con los totales correctos."""
        self.client.force_authenticate(user=self.user_a)
        response = self.client.get(
            reverse('kardex-report'),
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(response.status_code, 200)
        data = response.data

        # 3 movimientos esperados
        movs = data['movimientos']
        self.assertEqual(len(movs), 3)

        # Orden riguroso verificado: mov1, mov2, mov3
        self.assertEqual(movs[0]['id'], self.mov1.id)
        self.assertEqual(movs[1]['id'], self.mov2.id)
        self.assertEqual(movs[2]['id'], self.mov3.id)

        # Verificación Mov 1 (Entrada con Documento)
        self.assertEqual(movs[0]['documento_referencia'], '01 F001-00001234')
        self.assertEqual(Decimal(str(movs[0]['entrada_cantidad'])), Decimal('100.00000000'))
        self.assertEqual(Decimal(str(movs[0]['entrada_costo_unitario'])), Decimal('10.00000000'))
        self.assertEqual(Decimal(str(movs[0]['entrada_valor'])), Decimal('1000.00000000'))
        self.assertIsNone(movs[0]['salida_cantidad'])
        self.assertEqual(Decimal(str(movs[0]['saldo_cantidad'])), Decimal('100.00000000'))
        self.assertEqual(Decimal(str(movs[0]['cpp'])), Decimal('10.00000000'))

        # Verificación Mov 3 (Salida)
        self.assertIsNone(movs[2]['entrada_cantidad'])
        self.assertEqual(Decimal(str(movs[2]['salida_cantidad'])), Decimal('40.00000000'))
        self.assertEqual(Decimal(str(movs[2]['salida_costo_unitario'])), Decimal('11.00000000'))
        self.assertEqual(Decimal(str(movs[2]['salida_valor'])), Decimal('440.00000000'))
        self.assertEqual(Decimal(str(movs[2]['saldo_cantidad'])), Decimal('110.00000000'))
        self.assertEqual(Decimal(str(movs[2]['cpp'])), Decimal('11.00000000'))
        self.assertEqual(Decimal(str(movs[2]['saldo_valor'])), Decimal('1210.00000000'))

        # Verificación Resumen
        resumen = data['resumen']
        self.assertEqual(Decimal(str(resumen['total_entradas_cantidad'])), Decimal('150.00000000'))
        self.assertEqual(Decimal(str(resumen['total_entradas_valor'])), Decimal('1650.00000000'))
        self.assertEqual(Decimal(str(resumen['total_salidas_cantidad'])), Decimal('40.00000000'))
        self.assertEqual(Decimal(str(resumen['total_salidas_valor'])), Decimal('440.00000000'))
        self.assertEqual(Decimal(str(resumen['saldo_final_cantidad'])), Decimal('110.00000000'))
        self.assertEqual(Decimal(str(resumen['cpp_final'])), Decimal('11.00000000'))
        self.assertEqual(Decimal(str(resumen['saldo_final_valor'])), Decimal('1210.00000000'))

    def test_kardex_excel_export_coincidencia_1_a_1(self):
        """Verifica que el Excel generado tenga las 4 hojas y que Hoja 1 coincida 1:1 con la consulta en pantalla."""
        self.client.force_authenticate(user=self.user_a)
        response = self.client.get(
            reverse('kardex-excel-export'),
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(
            wb.sheetnames,
            ["Kardex Valorizado", "Movimientos", "Incidencias", "Resumen de Saldos"]
        )

        ws1 = wb["Kardex Valorizado"]
        # Filas 6 y 7 son encabezados, filas 8, 9, 10 son los 3 movimientos
        self.assertEqual(ws1.cell(row=8, column=1).value, "2026-01-10")
        self.assertEqual(ws1.cell(row=8, column=2).value, "01 F001-00001234")
        self.assertEqual(Decimal(str(ws1.cell(row=8, column=6).value)), Decimal('100.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=8, column=8).value)), Decimal('1000.0'))

        self.assertEqual(ws1.cell(row=9, column=1).value, "2026-01-15")
        self.assertEqual(Decimal(str(ws1.cell(row=9, column=6).value)), Decimal('50.0'))

        self.assertEqual(ws1.cell(row=10, column=1).value, "2026-01-20")
        self.assertEqual(Decimal(str(ws1.cell(row=10, column=9).value)), Decimal('40.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=10, column=11).value)), Decimal('440.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=10, column=12).value)), Decimal('110.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=10, column=13).value)), Decimal('11.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=10, column=14).value)), Decimal('1210.0'))

        # Fila 11: Fila de Totales
        self.assertEqual(ws1.cell(row=11, column=1).value, "TOTALES DEL REPORTE")
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=6).value)), Decimal('150.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=8).value)), Decimal('1650.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=9).value)), Decimal('40.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=11).value)), Decimal('440.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=12).value)), Decimal('110.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=13).value)), Decimal('11.0'))
        self.assertEqual(Decimal(str(ws1.cell(row=11, column=14).value)), Decimal('1210.0'))

    def test_aislamiento_multiempresa(self):
        """Verifica que Empresa A no pueda ver ni exportar el Kardex de Empresa B."""
        # User A consulta Empresa A -> sólo ve sus 3 movimientos
        self.client.force_authenticate(user=self.user_a)
        res_a = self.client.get(
            reverse('kardex-report'),
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(res_a.status_code, 200)
        self.assertEqual(len(res_a.data['movimientos']), 3)

        # User A intenta consultar con header de Empresa B -> 403 Forbidden
        res_prohibido = self.client.get(
            reverse('kardex-report'),
            HTTP_X_EMPRESA_ID=str(self.empresa_b.id)
        )
        self.assertEqual(res_prohibido.status_code, 403)

        # User A intenta exportar Excel con header de Empresa B -> 403 Forbidden
        res_excel_prohibido = self.client.get(
            reverse('kardex-excel-export'),
            HTTP_X_EMPRESA_ID=str(self.empresa_b.id)
        )
        self.assertEqual(res_excel_prohibido.status_code, 403)

        # User B consulta Empresa B -> sólo ve su movimiento
        self.client.force_authenticate(user=self.user_b)
        res_b = self.client.get(
            reverse('kardex-report'),
            HTTP_X_EMPRESA_ID=str(self.empresa_b.id)
        )
        self.assertEqual(res_b.status_code, 200)
        self.assertEqual(len(res_b.data['movimientos']), 1)
        self.assertEqual(res_b.data['movimientos'][0]['id'], self.mov_b.id)

    def test_filtros_kardex(self):
        """Verifica el filtrado por producto y rango de fechas."""
        self.client.force_authenticate(user=self.user_a)

        # Filtro por producto PROD-A
        res = self.client.get(
            reverse('kardex-report'),
            {'producto_id': self.prod_a.id},
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data['movimientos']), 3)

        # Filtro por rango de fechas (sólo del 12 al 18 de enero -> debe entrar únicamente mov2)
        res_fecha = self.client.get(
            reverse('kardex-report'),
            {'fecha_desde': '2026-01-12', 'fecha_hasta': '2026-01-18'},
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(res_fecha.status_code, 200)
        self.assertEqual(len(res_fecha.data['movimientos']), 1)
        self.assertEqual(res_fecha.data['movimientos'][0]['id'], self.mov2.id)

    def test_kardex_transferencias_saldo_progresivo(self):
        """Verifica que el saldo sea dinámico y progresivo tanto consolidado como por almacén."""
        self.client.force_authenticate(user=self.user_a)
        # Crear segundo almacén en Empresa A y un nuevo producto para prueba aislada
        almacen_a2 = Almacen.objects.create(
            empresa=self.empresa_a,
            establecimiento=self.estab_a,
            codigo='ALM02',
            nombre='ALMACÉN SECUNDARIO A'
        )
        prod_trans = Producto.objects.create(
            empresa=self.empresa_a,
            tipo_existencia_codigo=self.tipo_existencia,
            codigo_interno='PROD-TR',
            descripcion='PRODUCTO TRANSFERENCIAS'
        )
        op_tr_ent = TipoOperacion.objects.create(
            codigo_interno='TR_ENTRADA',
            codigo_sunat='21',
            descripcion='TRANSFERENCIA ENTRADA',
            naturaleza='ENTRADA',
            origen_requerido='TRANSFERENCIA',
        )
        op_tr_sal = TipoOperacion.objects.create(
            codigo_interno='TR_SALIDA',
            codigo_sunat='11',
            descripcion='TRANSFERENCIA SALIDA',
            naturaleza='SALIDA',
            origen_requerido='TRANSFERENCIA',
        )

        # 1. Saldo Inicial: 100 und en Almacén 1 (fecha 01/02)
        m1 = Movimiento.objects.create(
            empresa=self.empresa_a, producto=prod_trans, almacen=self.almacen_a1,
            periodo=self.periodo_a, tipo_operacion=self.op_compra,
            fecha_movimiento=date(2026, 2, 1), prioridad_dia=1, secuencia=1,
            naturaleza='ENTRADA', cantidad_base=Decimal('100.00000000'),
        )
        ValorizacionMovimiento.objects.create(
            movimiento=m1, costo_unitario_pen=Decimal('0.20000000'),
            valor_movimiento_pen=Decimal('20.00000000'),
            saldo_cantidad_empresa=Decimal('100.00000000'),
            saldo_valor_empresa=Decimal('20.00000000'),
            cpp_empresa=Decimal('0.20000000'),
        )

        # 2. Transferencia: 1 und entra a Almacén 2 (fecha 12/02, prioridad 2)
        m2 = Movimiento.objects.create(
            empresa=self.empresa_a, producto=prod_trans, almacen=almacen_a2,
            periodo=self.periodo_a, tipo_operacion=op_tr_ent,
            fecha_movimiento=date(2026, 2, 12), prioridad_dia=2, secuencia=1,
            naturaleza='ENTRADA', cantidad_base=Decimal('1.00000000'),
        )
        ValorizacionMovimiento.objects.create(
            movimiento=m2, costo_unitario_pen=Decimal('0.20000000'),
            valor_movimiento_pen=Decimal('0.20000000'),
            saldo_cantidad_empresa=Decimal('100.00000000'),
            saldo_valor_empresa=Decimal('20.00000000'),
            cpp_empresa=Decimal('0.20000000'),
        )

        # 3. Transferencia: 1 und sale de Almacén 1 (fecha 12/02, prioridad 4)
        m3 = Movimiento.objects.create(
            empresa=self.empresa_a, producto=prod_trans, almacen=self.almacen_a1,
            periodo=self.periodo_a, tipo_operacion=op_tr_sal,
            fecha_movimiento=date(2026, 2, 12), prioridad_dia=4, secuencia=1,
            naturaleza='SALIDA', cantidad_base=Decimal('1.00000000'),
        )
        ValorizacionMovimiento.objects.create(
            movimiento=m3, costo_unitario_pen=Decimal('0.20000000'),
            valor_movimiento_pen=Decimal('0.20000000'),
            saldo_cantidad_empresa=Decimal('100.00000000'),
            saldo_valor_empresa=Decimal('20.00000000'),
            cpp_empresa=Decimal('0.20000000'),
        )

        # 4. Venta: 5 und salen de Almacén 1 (fecha 22/02, prioridad 4)
        m4 = Movimiento.objects.create(
            empresa=self.empresa_a, producto=prod_trans, almacen=self.almacen_a1,
            periodo=self.periodo_a, tipo_operacion=self.op_venta,
            fecha_movimiento=date(2026, 2, 22), prioridad_dia=4, secuencia=1,
            naturaleza='SALIDA', cantidad_base=Decimal('5.00000000'),
        )
        ValorizacionMovimiento.objects.create(
            movimiento=m4, costo_unitario_pen=Decimal('0.20000000'),
            valor_movimiento_pen=Decimal('1.00000000'),
            saldo_cantidad_empresa=Decimal('95.00000000'),
            saldo_valor_empresa=Decimal('19.00000000'),
            cpp_empresa=Decimal('0.20000000'),
        )

        # A) Consulta Consolidada (Sin filtro de almacén)
        res_global = self.client.get(
            reverse('kardex-report'),
            {'producto_id': prod_trans.id},
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(res_global.status_code, 200)
        movs_g = res_global.data['movimientos']
        self.assertEqual(len(movs_g), 4)
        # Fila 1: Saldo 100
        self.assertEqual(Decimal(str(movs_g[0]['saldo_cantidad'])), Decimal('100.00000000'))
        # Fila 2: Entra 1 -> Saldo 101
        self.assertEqual(Decimal(str(movs_g[1]['saldo_cantidad'])), Decimal('101.00000000'))
        # Fila 3: Sale 1 -> Saldo 100
        self.assertEqual(Decimal(str(movs_g[2]['saldo_cantidad'])), Decimal('100.00000000'))
        # Fila 4: Sale 5 -> Saldo 95
        self.assertEqual(Decimal(str(movs_g[3]['saldo_cantidad'])), Decimal('95.00000000'))
        # Totales Consolidados
        resumen_g = res_global.data['resumen']
        self.assertEqual(Decimal(str(resumen_g['total_entradas_cantidad'])), Decimal('101.00000000'))
        self.assertEqual(Decimal(str(resumen_g['total_salidas_cantidad'])), Decimal('6.00000000'))
        self.assertEqual(Decimal(str(resumen_g['saldo_final_cantidad'])), Decimal('95.00000000'))

        # B) Consulta Filtrada por Almacén 1
        res_alm1 = self.client.get(
            reverse('kardex-report'),
            {'producto_id': prod_trans.id, 'almacen_id': self.almacen_a1.id},
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(res_alm1.status_code, 200)
        movs_1 = res_alm1.data['movimientos']
        self.assertEqual(len(movs_1), 3)
        # Fila 1: Entra 100 -> Saldo 100
        self.assertEqual(Decimal(str(movs_1[0]['saldo_cantidad'])), Decimal('100.00000000'))
        # Fila 2: Sale 1 (transferencia) -> Saldo 99
        self.assertEqual(Decimal(str(movs_1[1]['saldo_cantidad'])), Decimal('99.00000000'))
        # Fila 3: Sale 5 (venta) -> Saldo 94
        self.assertEqual(Decimal(str(movs_1[2]['saldo_cantidad'])), Decimal('94.00000000'))
        # Totales Almacén 1: 100 - 6 = 94
        resumen_1 = res_alm1.data['resumen']
        self.assertEqual(Decimal(str(resumen_1['total_entradas_cantidad'])), Decimal('100.00000000'))
        self.assertEqual(Decimal(str(resumen_1['total_salidas_cantidad'])), Decimal('6.00000000'))
        self.assertEqual(Decimal(str(resumen_1['saldo_final_cantidad'])), Decimal('94.00000000'))

        # C) Consulta Filtrada por Almacén 2
        res_alm2 = self.client.get(
            reverse('kardex-report'),
            {'producto_id': prod_trans.id, 'almacen_id': almacen_a2.id},
            HTTP_X_EMPRESA_ID=str(self.empresa_a.id)
        )
        self.assertEqual(res_alm2.status_code, 200)
        movs_2 = res_alm2.data['movimientos']
        self.assertEqual(len(movs_2), 1)
        # Fila 1: Entra 1 (transferencia) -> Saldo 1
        self.assertEqual(Decimal(str(movs_2[0]['saldo_cantidad'])), Decimal('1.00000000'))
        resumen_2 = res_alm2.data['resumen']
        self.assertEqual(Decimal(str(resumen_2['total_entradas_cantidad'])), Decimal('1.00000000'))
        self.assertEqual(Decimal(str(resumen_2['total_salidas_cantidad'])), Decimal('0.00000000'))
        self.assertEqual(Decimal(str(resumen_2['saldo_final_cantidad'])), Decimal('1.00000000'))
