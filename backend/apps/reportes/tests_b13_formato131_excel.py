"""
Bloque 13 — exportación oficial Excel del Formato 13.1.

La descarga fallaba "en silencio" en el frontend (cuerpo de error como Blob) y la URL del Blob
se liberaba en el mismo click; el backend respondía bien. Aquí se fija el contrato del archivo:
HTTP y cabeceras, XLSX legible, mismos valores que la pantalla (JSON del 13.1) y el cierre,
microdiferencia, nombre seguro, inyección de fórmulas, multiempresa y versiones de cierre.
"""
import io
import re
import zipfile
from decimal import Decimal

import openpyxl
from django.test import TestCase
from django.urls import reverse

from apps.documentos.fixtures_bloque2 import Bloque2Fixture
from apps.empresas.models import CierrePeriodo, Establecimiento
from apps.inventario.models import Producto, ProductoUnidad

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
D8 = Decimal('0.00000001')


def d(valor):
    return Decimal(str(valor)).quantize(D8)


class Formato131ExcelBase(Bloque2Fixture):
    def setUp(self):
        self.crear_base()
        # Saldo inicial con microdiferencia exacta: 120 u por S/ 1200.00004 → CPP 10.000000
        # (6 decimales) y 120 × 10.000000 = 1200.00000 → microdiferencia oficial 0.00004000.
        self.saldo_inicial(self.producto, '120', '1200.00004')
        self.compra_confirmada('10', '10', fecha='2026-01-10')

    def saldo_inicial(self, producto, cantidad, valor):
        resp = self.client.post(reverse('saldo-inicial-list-create'), {
            'periodo': self.enero.id, 'almacen': self.almacen.id, 'producto': producto.id, 'cantidad': cantidad, 'valor_pen': valor,
        }, format='json', **self.headers)
        self.assertEqual(resp.status_code, 201, resp.data)
        conf = self.client.post(reverse('saldo-inicial-confirmar', args=[resp.data['id']]), format='json', **self.headers)
        self.assertEqual(conf.status_code, 200, conf.data)

    def cerrar_enero(self):
        resp = self.client.post(reverse('periodo-cerrar', args=[self.enero.id]), format='json', **self.headers)
        self.assertEqual(resp.status_code, 200, resp.data)

    def exportar(self, **params):
        params.setdefault('periodo_id', self.enero.id)
        return self.client.get(reverse('formato-131-excel'), params, **self.headers)

    def pantalla(self, **params):
        params.setdefault('periodo_id', self.enero.id)
        resp = self.client.get(reverse('formato-131'), params, **self.headers)
        self.assertEqual(resp.status_code, 200, resp.data)
        return resp.data

    def libro(self, resp):
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    def hoja_de(self, libro, codigo):
        return next(h for h in libro.worksheets if h['B6'].value == codigo)


class ExportacionFormato131Tests(Formato131ExcelBase, TestCase):
    def test_http_cabeceras_y_archivo_valido(self):
        self.cerrar_enero()
        resp = self.exportar(producto_id=self.producto.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], XLSX)
        self.assertEqual(resp['Content-Disposition'], 'attachment; filename="Formato_13_1_2026-01_20100100100.xlsx"')
        self.assertEqual(int(resp['Content-Length']), len(resp.content))
        self.assertGreater(len(resp.content), 1000)
        self.assertEqual(resp.content[:2], b'PK')
        self.assertIsNone(zipfile.ZipFile(io.BytesIO(resp.content)).testzip())  # no está corrupto
        libro = self.libro(resp)
        self.assertEqual(libro.sheetnames, ['13.1 01'])
        hoja = libro.active
        self.assertEqual(hoja['A1'].value, 'FORMATO 13.1: REGISTRO DE INVENTARIO PERMANENTE VALORIZADO')
        self.assertEqual([hoja.cell(r, 2).value for r in (2, 3, 4)], ['2026-01', '20100100100', 'EMPRESA B2 S.A.C.'])
        self.assertEqual(hoja['A13'].value, 'SALDO INICIAL')
        self.assertEqual([hoja.cell(12, c).value for c in range(1, 15)][:6], ['Fecha', 'Tipo', 'Serie', 'Número', 'Código Tabla 12', 'Cantidad'])

    def test_excel_igual_a_la_pantalla_y_al_cierre_con_microdiferencia(self):
        self.cerrar_enero()
        datos = self.pantalla(producto_id=self.producto.id)
        reporte = datos['reportes'][0]
        self.assertEqual(reporte['microdiferencia_cierre'], '0.00004000')
        self.assertEqual(reporte['fuente_saldo_final'], 'CIERRE')
        hoja = self.libro(self.exportar(producto_id=self.producto.id)).active
        # Una fila por movimiento, con los mismos importes de la pantalla (8 decimales).
        for i, mov in enumerate(reporte['movimientos'], start=14):
            self.assertEqual(hoja.cell(i, 1).value, mov['fecha_emision'])
            for col, bloque in ((6, mov['entrada']), (9, mov['salida']), (12, mov['saldo_final'])):
                for k, clave in enumerate(('cantidad', 'costo_unitario', 'costo_total')):
                    esperado = d(bloque[clave]) if bloque else None
                    valor = hoja.cell(i, col + k).value
                    self.assertEqual(d(valor) if valor is not None else None, esperado, (i, col + k))
        ultima = 13 + len(reporte['movimientos'])
        final = reporte['saldo_final']
        # F131-001: después de los movimientos, la fila explícita de saldo final del período.
        self.assertEqual(hoja.cell(ultima + 1, 1).value, 'SALDO FINAL DEL PERÍODO')
        self.assertEqual([d(hoja.cell(ultima + 1, c).value) for c in (12, 13, 14)],
                         [d(final['cantidad']), d(final['costo_unitario']), d(final['costo_total'])])
        # Cierre / snapshot: el mismo saldo.
        cierre = CierrePeriodo.objects.get(periodo=self.enero, vigente=True)
        detalle = cierre.detalles.get(producto=self.producto)
        self.assertEqual((d(detalle.saldo_cantidad), d(detalle.saldo_valor_pen)), (d(final['cantidad']), d(final['costo_total'])))
        self.assertEqual((d(final['cantidad']), d(final['costo_total'])), (Decimal('130'), Decimal('1300.00004')))
        # Pie: fuente del saldo y microdiferencia absorbida, igual que el aviso de la pantalla.
        self.assertEqual(hoja.cell(ultima + 3, 1).value, 'FUENTE DEL SALDO FINAL')
        self.assertTrue(hoja.cell(ultima + 3, 2).value.startswith(f'CIERRE v1 (hash {cierre.hash_resultado[:12]}'))
        self.assertIn('MICRODIFERENCIA OFICIAL', hoja.cell(ultima + 4, 1).value)
        self.assertEqual(d(hoja.cell(ultima + 4, 14).value), Decimal('0.00004000'))
        self.assertEqual(hoja.print_area, f"'13.1 01'!$A$1:$N${ultima + 4}")

    def test_la_precision_no_se_pierde_en_el_archivo(self):
        # Cada número del XML, leído como texto, es exactamente el importe de 8 decimales.
        self.cerrar_enero()
        resp = self.exportar(producto_id=self.producto.id)
        hoja_xml = zipfile.ZipFile(io.BytesIO(resp.content)).read('xl/worksheets/sheet1.xml').decode()
        numeros = {d(Decimal(v)) for v in re.findall(r'<v>([^<]+)</v>', hoja_xml)}
        for esperado in ('1200.00004', '1300.00004', '0.00004'):
            self.assertIn(Decimal(esperado), numeros)

    def test_periodo_abierto_se_exporta_como_calculado(self):
        hoja = self.libro(self.exportar(producto_id=self.producto.id)).active
        filas = [hoja.cell(r, 1).value for r in range(1, hoja.max_row + 1)]
        fila = filas.index('FUENTE DEL SALDO FINAL') + 1
        self.assertEqual(hoja.cell(fila, 2).value, 'CALCULADO (sin cierre vigente)')

    def test_sin_movimientos_ni_saldos_no_genera_un_excel_invalido(self):
        nuevo = Producto.objects.create(empresa=self.empresa, codigo_interno='SIN-MOV', descripcion='SIN MOV',
                                        tipo_existencia_codigo=self.tipo_existencia)
        ProductoUnidad.objects.create(producto=nuevo, unidad_codigo=self.und, factor_a_base=Decimal('1'), es_base=True)
        resp = self.exportar(producto_id=nuevo.id)
        self.assertEqual(resp.status_code, 200)
        libro = self.libro(resp)
        self.assertEqual(libro.sheetnames, ['Formato 13.1'])
        self.assertEqual(libro.active['A1'].value, 'No existen movimientos para los filtros seleccionados.')


class VersionesDeCierreTests(Formato131ExcelBase, TestCase):
    def test_exporta_el_cierre_vigente_v2_sin_mezclar_v1(self):
        self.cerrar_enero()
        v1 = CierrePeriodo.objects.get(periodo=self.enero, vigente=True)
        reabrir = self.client.post(reverse('periodo-reabrir', args=[self.enero.id]), {'motivo_reapertura': 'QA bloque 13'}, format='json', **self.headers)
        self.assertEqual(reabrir.status_code, 200, reabrir.data)
        self.compra_confirmada('10', '20', fecha='2026-01-20')
        self.cerrar_enero()
        v2 = CierrePeriodo.objects.get(periodo=self.enero, vigente=True)
        self.assertEqual((v1.version_calculo, v2.version_calculo), (1, 2))
        hoja = self.libro(self.exportar(producto_id=self.producto.id)).active
        filas = [hoja.cell(r, 1).value for r in range(1, hoja.max_row + 1)]
        fila_fuente = filas.index('FUENTE DEL SALDO FINAL') + 1
        self.assertTrue(hoja.cell(fila_fuente, 2).value.startswith(f'CIERRE v2 (hash {v2.hash_resultado[:12]}'))
        ultima = fila_fuente - 2
        final = v2.detalles.get(producto=self.producto)
        self.assertEqual(d(hoja.cell(ultima, 14).value), d(final.saldo_valor_pen))
        self.assertNotEqual(d(final.saldo_valor_pen), d(v1.detalles.get(producto=self.producto).saldo_valor_pen))


class SeguridadExportacionTests(Formato131ExcelBase, TestCase):
    def test_textos_con_formula_quedan_como_texto(self):
        for i, texto in enumerate(('=HYPERLINK("http://evil","x")', '+SUM(A1:A2)', '-2+3', '@cmd')):
            p = Producto.objects.create(empresa=self.empresa, codigo_interno=texto[:20] if i else '=1+1',
                                        descripcion=texto, tipo_existencia_codigo=self.tipo_existencia)
            ProductoUnidad.objects.create(producto=p, unidad_codigo=self.und, factor_a_base=Decimal('1'), es_base=True)
            self.documento_confirmado('COMPRA', [self.linea('COMPRA', '1', '10', producto=p)], fecha='2026-01-11')
        Establecimiento.objects.filter(pk=self.establecimiento.pk).update(nombre='=cmd|"/c calc"!A1')
        resp = self.exportar()
        self.assertEqual(resp.status_code, 200)
        archivo = zipfile.ZipFile(io.BytesIO(resp.content))
        hojas_xml = [archivo.read(n).decode() for n in archivo.namelist() if n.startswith('xl/worksheets/sheet')]
        self.assertFalse(any('<f>' in x for x in hojas_xml), 'no debe haber ninguna fórmula en el archivo')
        libro = self.libro(resp)
        textos = {c.value for h in libro.worksheets for f in h.iter_rows() for c in f if isinstance(c.value, str)}
        for esperado in ('=HYPERLINK("http://evil","x")', '+SUM(A1:A2)', '-2+3', '@cmd', '=1+1'):
            self.assertIn(esperado, textos)
        self.assertTrue(any('=cmd|' in t for t in textos))
        self.assertFalse(any(c.data_type == 'f' for h in libro.worksheets for f in h.iter_rows() for c in f))

    def test_no_exporta_periodo_de_otra_empresa(self):
        from apps.empresas.models import Periodo
        ajeno = Periodo.objects.create(empresa=self.otra_empresa, anio=2026, mes=1, orden=202601,
                                       fecha_inicio=self.enero.fecha_inicio, fecha_fin=self.enero.fecha_fin, estado='CERRADO')
        resp = self.exportar(periodo_id=ajeno.id)
        self.assertEqual(resp.status_code, 400)
        self.assertNotEqual(resp.get('Content-Type'), XLSX)
        self.assertIn('no pertenece a la empresa activa', str(resp.data))

    def test_no_exporta_producto_ni_establecimiento_de_otra_empresa(self):
        ajeno = Producto.objects.create(empresa=self.otra_empresa, codigo_interno='AJ', descripcion='AJENO', tipo_existencia_codigo=self.tipo_existencia)
        est = Establecimiento.objects.create(empresa=self.otra_empresa, codigo_sunat='0001', nombre='AJENO')
        self.assertEqual(self.exportar(producto_id=ajeno.id).status_code, 400)
        self.assertEqual(self.exportar(establecimiento_id=est.id).status_code, 400)

    def test_no_exporta_con_la_cabecera_de_otra_empresa(self):
        resp = self.client.get(reverse('formato-131-excel'), {'periodo_id': self.enero.id}, HTTP_X_EMPRESA_ID=str(self.otra_empresa.id))
        self.assertEqual(resp.status_code, 403)
        self.assertNotEqual(resp.get('Content-Type'), XLSX)

    def test_periodo_inexistente_o_sin_parametro(self):
        for params in ({'periodo_id': 999999}, {'periodo_id': ''}, {'periodo_id': 'x'}):
            with self.subTest(params=params):
                resp = self.client.get(reverse('formato-131-excel'), params, **self.headers)
                self.assertEqual(resp.status_code, 400)
                self.assertNotEqual(resp.get('Content-Type'), XLSX)
                self.assertTrue(re.search('periodo_id', str(resp.data)))
