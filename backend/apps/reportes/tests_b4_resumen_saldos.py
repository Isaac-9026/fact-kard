"""
Bloque 4 — la hoja "Resumen de Saldos" del Excel del Kardex usaba StockAlmacen y
SaldoCppActual (estado de HOY): un Excel de enero cambiaba tras una compra de febrero.
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.test import TestCase
from django.urls import reverse

from apps.documentos.fixtures_bloque2 import Bloque2Fixture
from apps.empresas.models import CierrePeriodoDetalle, Periodo
from apps.empresas.services import cerrar_periodo
from apps.inventario.models import Almacen, Transferencia, TransferenciaDetalle
from apps.inventario.services import confirmar_transferencia
from apps.catalogos.models import TipoOperacion
from apps.reportes.formato131_service import generar_formato_131
from apps.reportes.saldos_corte import repartir_valor
from apps.reportes.services import consultar_kardex_valorizado

D = Decimal


class ResumenSaldosHistoricoTests(Bloque2Fixture, TestCase):
    def setUp(self):
        self.crear_base()
        self.febrero = Periodo.objects.create(
            empresa=self.empresa, anio=2026, mes=2, orden=202602,
            fecha_inicio=date(2026, 2, 1), fecha_fin=date(2026, 2, 28), estado='ABIERTO',
        )
        TipoOperacion.objects.create(codigo_interno='TRANSFERENCIA_SALIDA', codigo_sunat='11', descripcion='TRANSF. SALIDA',
                                     naturaleza='SALIDA', origen_requerido='TRANSFERENCIA', es_interna=True)
        TipoOperacion.objects.create(codigo_interno='TRANSFERENCIA_ENTRADA', codigo_sunat='21', descripcion='TRANSF. ENTRADA',
                                     naturaleza='ENTRADA', origen_requerido='TRANSFERENCIA', es_interna=True)
        self.almacen_2 = Almacen.objects.create(empresa=self.empresa, establecimiento=self.establecimiento,
                                                codigo='ALM2', nombre='SEGUNDO')
        # Enero: 3 u. por S/ 2.00 (CPP 0.666667), 1 u. pasa al segundo almacén, se vende 1.
        self.compra_confirmada('1', '1', fecha='2026-01-08')
        self.compra_confirmada('2', '0.5', fecha='2026-01-09')
        transferencia = Transferencia.objects.create(empresa=self.empresa, almacen_origen=self.almacen,
                                                     almacen_destino=self.almacen_2, fecha=date(2026, 1, 10),
                                                     creado_por=self.usuario)
        TransferenciaDetalle.objects.create(transferencia=transferencia, producto=self.producto, cantidad_base=D('1'))
        confirmar_transferencia(transferencia_id=transferencia.id, empresa_id=self.empresa.id, usuario=self.usuario)
        self.documento_confirmado('VENTA', [self.linea('VENTA', '1', '30')], fecha='2026-01-20')

    def resumen_excel(self, **filtros):
        resp = self.client.get(reverse('kardex-excel-export'), filtros, **self.headers)
        self.assertEqual(resp.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(resp.content))['Resumen de Saldos']
        filas = []
        for fila in ws.iter_rows(min_row=4, values_only=True):
            if not fila[0]:
                break
            filas.append(fila[:9])
        return ws['A2'].value, filas

    def compra_febrero(self):
        self.documento_confirmado('COMPRA', [self.linea('COMPRA', '7', '13')], fecha='2026-02-03', periodo=self.febrero)

    def assert_concilia(self, filas):
        # Total empresa (2 u. / 1.333333) = Kardex = 13.1 = suma por almacén.
        kardex = consultar_kardex_valorizado(empresa_id=self.empresa.id, producto_id=self.producto.id,
                                             periodo_id=self.enero.id)['resumen']
        f131 = [r for r in generar_formato_131(empresa_id=self.empresa.id, periodo_id=self.enero.id)['reportes']
                if r['cabecera']['producto_id'] == self.producto.id]
        cantidad_empresa, valor_empresa = D(str(filas[0][7])), D(str(filas[0][8]))
        self.assertEqual((cantidad_empresa, valor_empresa), (D('2'), D('1.333333')))
        self.assertEqual(kardex['saldo_final_valor'], valor_empresa)
        self.assertEqual(sum(D(r['saldo_final']['costo_total']) for r in f131), valor_empresa)
        self.assertEqual(sum(D(str(f[6])) for f in filas), valor_empresa)  # sin crear ni perder importes
        self.assertEqual(sum(D(str(f[4])) for f in filas), cantidad_empresa)

    # 1-4 período abierto: saldo histórico reconstruido
    def test_enero_abierto_no_cambia_tras_compra_de_febrero(self):
        antes = self.resumen_excel(periodo_id=self.enero.id)
        self.assertIn('HISTORICO al 2026-01-31', antes[0])
        self.assertEqual([f[3] for f in antes[1]], ['ALM - ALMACEN', 'ALM2 - SEGUNDO'])
        self.compra_febrero()
        despues = self.resumen_excel(periodo_id=self.enero.id)
        self.assertEqual(antes, despues)
        self.assert_concilia(despues[1])
        # El estado actual sí refleja febrero.
        _, actual = self.resumen_excel()
        self.assertEqual(D(str(actual[0][7])), D('9'))

    def test_fecha_hasta_corta_el_resumen(self):
        _, filas = self.resumen_excel(fecha_hasta='2026-01-09')
        self.assertEqual(len(filas), 1)
        self.assertEqual((D(str(filas[0][4])), D(str(filas[0][8]))), (D('3'), D('2')))

    # 5 período cerrado: snapshot del cierre
    def test_enero_cerrado_usa_el_cierre_y_concilia(self):
        cierre, _ = cerrar_periodo(empresa_id=self.empresa.id, periodo_id=self.enero.id, usuario=self.usuario)
        self.compra_febrero()
        cabecera, filas = self.resumen_excel(periodo_id=self.enero.id)
        self.assertIn(f'CIERRE v{cierre.version_calculo}', cabecera)
        detalle = CierrePeriodoDetalle.objects.get(cierre_periodo=cierre, producto=self.producto)
        self.assertEqual((D(str(filas[0][7])), D(str(filas[0][8]))), (detalle.saldo_cantidad, detalle.saldo_valor_pen))
        self.assert_concilia(filas)
        self.assertEqual(self.resumen_excel(periodo_id=self.enero.id), (cabecera, filas))


class RepartirValorTests(TestCase):
    def test_no_crea_ni_pierde_importes(self):
        valores = repartir_valor(D('2.00000000'), [(1, D('1')), (2, D('1')), (3, D('1'))], D('0.666667'))
        self.assertEqual(sum(valores.values()), D('2.00000000'))
        self.assertEqual(valores, {1: D('0.666667'), 2: D('0.666667'), 3: D('0.666666')})

    def test_diferencia_va_a_la_parte_preferida(self):
        valores = repartir_valor(D('2'), [(1, D('1')), (2, D('2'))], D('0.666667'), preferida=1)
        self.assertEqual(valores, {1: D('0.666666'), 2: D('1.333334')})
