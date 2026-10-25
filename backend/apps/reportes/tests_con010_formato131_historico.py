"""
Regresión CON-010: el Formato 13.1 usaba SaldoCppActual (CPP de HOY) para productos
sin movimientos en el período y no valorizaba el saldo inicial; un período pasado
cambiaba al registrar operaciones posteriores.
"""
from datetime import date
from decimal import Decimal

from django.test import TestCase

from apps.documentos.fixtures_bloque2 import Bloque2Fixture
from apps.empresas.models import CierrePeriodoDetalle, Periodo
from apps.empresas.services import cerrar_periodo
from apps.reportes.formato131_excel import generar_excel_formato_131
from apps.reportes.formato131_service import generar_formato_131
from apps.reportes.services import consultar_kardex_valorizado


class Formato131HistoricoTests(Bloque2Fixture, TestCase):
    def setUp(self):
        self.crear_base()
        Periodo.objects.filter(pk=self.diciembre.pk).update(estado='ABIERTO')
        self.febrero = Periodo.objects.create(
            empresa=self.empresa, anio=2026, mes=2, orden=202602,
            fecha_inicio=date(2026, 2, 1), fecha_fin=date(2026, 2, 28), estado='ABIERTO',
        )
        # Diciembre: producto_2 entra 5 × 8.00 y no vuelve a moverse en enero.
        self.compra_confirmada_de(self.producto_2, '5', '8', '2025-12-15', self.diciembre)
        # Enero: producto entra 10 × 10.00 y se venden 4.
        self.compra_confirmada_de(self.producto, '10', '10', '2026-01-10', self.enero)
        self.documento_confirmado('VENTA', [self.linea('VENTA', '4', '30')], fecha='2026-01-20', periodo=self.enero)

    def compra_confirmada_de(self, producto, cantidad, precio, fecha, periodo):
        return self.documento_confirmado(
            'COMPRA', [self.linea('COMPRA', cantidad, precio, producto=producto)], fecha=fecha, periodo=periodo,
        )

    def reporte(self, periodo, producto):
        data = generar_formato_131(empresa_id=self.empresa.id, periodo_id=periodo.id)
        return next(r for r in data['reportes'] if r['cabecera']['producto_id'] == producto.id)

    def _operaciones_posteriores_con_otro_cpp(self):
        self.compra_confirmada_de(self.producto_2, '5', '20', '2026-02-05', self.febrero)
        self.compra_confirmada_de(self.producto, '10', '30', '2026-02-05', self.febrero)

    def test_enero_no_cambia_tras_compras_de_febrero(self):
        antes = generar_formato_131(empresa_id=self.empresa.id, periodo_id=self.enero.id)

        self._operaciones_posteriores_con_otro_cpp()

        despues = generar_formato_131(empresa_id=self.empresa.id, periodo_id=self.enero.id)
        self.assertEqual(antes, despues)

    def test_producto_sin_movimientos_en_el_periodo_usa_cpp_historico(self):
        self._operaciones_posteriores_con_otro_cpp()
        r = self.reporte(self.enero, self.producto_2)
        self.assertEqual(r['movimientos'], [])
        self.assertEqual(r['saldo_inicial_valorizado'], {
            'cantidad': '5.00000000', 'costo_unitario': '8.00000000', 'costo_total': '40.00000000',
        })
        self.assertEqual(r['saldo_final'], {
            'cantidad': '5.00000000', 'costo_unitario': '8.00000000', 'costo_total': '40.00000000',
        })

    def test_saldo_inicial_y_final_valorizados(self):
        self._operaciones_posteriores_con_otro_cpp()
        r = self.reporte(self.febrero, self.producto)
        self.assertEqual(r['saldo_inicial_valorizado'], {
            'cantidad': '6.00000000', 'costo_unitario': '10.00000000', 'costo_total': '60.00000000',
        })
        self.assertEqual(r['saldo_inicial'], '6.00000000')  # compatibilidad con el campo anterior
        # 6 × 10 + 10 × 30 = 360 → CPP 22.50
        self.assertEqual(r['saldo_final'], {
            'cantidad': '16.00000000', 'costo_unitario': '22.50000000', 'costo_total': '360.00000000',
        })

    def test_kardex_igual_formato_131_igual_cierre(self):
        cerrar_periodo(empresa_id=self.empresa.id, periodo_id=self.diciembre.id, usuario=self.usuario)
        cierre, _ = cerrar_periodo(empresa_id=self.empresa.id, periodo_id=self.enero.id, usuario=self.usuario)
        self._operaciones_posteriores_con_otro_cpp()

        for producto, esperado in ((self.producto, ('6', '60')), (self.producto_2, ('5', '40'))):
            with self.subTest(producto=producto.codigo_interno):
                detalle = CierrePeriodoDetalle.objects.get(cierre_periodo=cierre, producto=producto)
                f131 = self.reporte(self.enero, producto)['saldo_final']
                kardex = consultar_kardex_valorizado(
                    empresa_id=self.empresa.id, producto_id=producto.id, fecha_hasta='2026-01-31',
                )['resumen']
                cantidad, valor = (Decimal(v) for v in esperado)
                self.assertEqual(detalle.saldo_cantidad, cantidad)
                self.assertEqual(detalle.saldo_valor_pen, valor)
                self.assertEqual(Decimal(f131['cantidad']), cantidad)
                self.assertEqual(Decimal(f131['costo_total']), valor)
                self.assertEqual(kardex['saldo_final_cantidad'], cantidad)
                self.assertEqual(kardex['saldo_final_valor'], valor)

        # Febrero arranca exactamente donde terminó el cierre de enero.
        inicial = self.reporte(self.febrero, self.producto_2)['saldo_inicial_valorizado']
        self.assertEqual(inicial['costo_total'], '40.00000000')

    def test_excel_muestra_saldo_inicial_valorizado(self):
        import io
        from openpyxl import load_workbook
        self._operaciones_posteriores_con_otro_cpp()
        libro = load_workbook(io.BytesIO(
            generar_excel_formato_131(empresa_id=self.empresa.id, periodo_id=self.febrero.id).getvalue()
        ))
        hoja = next(h for h in libro.worksheets if h['B6'].value == 'P-B2')
        self.assertEqual(hoja['A13'].value, 'SALDO INICIAL')
        self.assertEqual([hoja.cell(13, c).value for c in (12, 13, 14)],
                         [Decimal('6'), Decimal('10'), Decimal('60')])
        self.assertEqual(hoja['A14'].value, '2026-02-05')
