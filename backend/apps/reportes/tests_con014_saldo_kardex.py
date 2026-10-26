"""
CON-014 — el saldo valorizado del Kardex era cantidad × CPP (redondeado) y el saldo
progresivo arrancaba en el primer movimiento visible. Ahora cada fila muestra el valor
contable real acumulado por el motor y existe un saldo de arranque para el filtro.
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.test import TestCase
from django.urls import reverse

from apps.catalogos.models import TipoOperacion
from apps.documentos.fixtures_bloque2 import Bloque2Fixture
from apps.documentos.models import Documento, DocumentoDetalle
from apps.inventario.models import Almacen, Transferencia, TransferenciaDetalle
from apps.inventario.services import confirmar_transferencia
from apps.kardex.models import Movimiento, ValorizacionMovimiento

ORDEN = ('fecha_movimiento', 'prioridad_dia', 'secuencia', 'id')
D = Decimal


class SaldoKardexBase(Bloque2Fixture):
    def setUp(self):
        self.crear_base()

    def kardex(self, **filtros):
        filtros.setdefault('producto_id', self.producto.id)
        resp = self.client.get(reverse('kardex-report'), filtros, **self.headers)
        self.assertEqual(resp.status_code, 200, getattr(resp, 'data', None))
        return resp.data

    def excel(self, **filtros):
        filtros.setdefault('producto_id', self.producto.id)
        resp = self.client.get(reverse('kardex-excel-export'), filtros, **self.headers)
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content))['Kardex Valorizado']

    def venta(self, cantidad, fecha='2026-01-12'):
        return self.documento_confirmado('VENTA', [self.linea('VENTA', cantidad, '30')], fecha=fecha)

    def ultima_valorizacion(self, producto=None):
        mov = self.movimientos(producto).order_by(*ORDEN).last()
        return ValorizacionMovimiento.objects.get(movimiento=mov)

    def assert_filas_contables(self, data):
        """Cada fila = valorización real del motor y continuidad saldo + efecto = nuevo."""
        anterior_valor = D(data['resumen']['saldo_inicial_valor'])
        anterior_cant = D(data['resumen']['saldo_inicial_cantidad'])
        for fila in data['movimientos']:
            val = ValorizacionMovimiento.objects.get(movimiento_id=fila['id'])
            self.assertEqual(D(fila['saldo_valor']), val.saldo_valor_empresa)
            self.assertEqual(D(fila['saldo_cantidad']), val.saldo_cantidad_empresa)
            self.assertEqual(D(fila['cpp']), val.cpp_empresa)
            efecto = D(fila['entrada_valor'] or 0) - D(fila['salida_valor'] or 0)
            mov = Movimiento.objects.get(pk=fila['id'])
            if mov.transferencia_id:
                efecto = D('0')  # redistribuye entre almacenes; no altera el saldo de empresa
            self.assertEqual(anterior_valor + efecto, D(fila['saldo_valor']), fila)
            cant = D(fila['entrada_cantidad'] or 0) - D(fila['salida_cantidad'] or 0)
            if not mov.transferencia_id:
                self.assertEqual(anterior_cant + cant, D(fila['saldo_cantidad']), fila)
            anterior_valor, anterior_cant = D(fila['saldo_valor']), D(fila['saldo_cantidad'])
        self.assertEqual(anterior_valor, D(data['resumen']['saldo_final_valor']))
        self.assertEqual(anterior_cant, D(data['resumen']['saldo_final_cantidad']))


class SaldoKardexTests(SaldoKardexBase, TestCase):

    # 1
    def test_compra_normal(self):
        self.compra_confirmada('10', '10')
        data = self.kardex()
        self.assertEqual(D(data['movimientos'][0]['saldo_valor']), D('100'))
        self.assert_filas_contables(data)

    # 2
    def test_venta(self):
        self.compra_confirmada('10', '10')
        self.venta('4')
        data = self.kardex()
        self.assertEqual(D(data['movimientos'][-1]['saldo_valor']), D('60'))
        self.assert_filas_contables(data)

    def test_redondeo_no_se_reconstruye_con_cantidad_por_cpp(self):
        # 1 × 1.00 + 2 × 0.50 = 2.00 por 3 u. → CPP 0.666667 (6 decimales); 3 × CPP = 2.00000100.
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        fila = self.kardex()['movimientos'][-1]
        self.assertEqual(D(fila['cpp']), D('0.66666700'))
        self.assertEqual(D(fila['saldo_valor']), D('2.00000000'))
        self.assertNotEqual(D(fila['saldo_valor']), D(fila['saldo_cantidad']) * D(fila['cpp']))

    # 3
    def test_ajuste_de_valor(self):
        compra = self.compra_confirmada('3', '1')
        detalle = DocumentoDetalle.objects.get(documento_id=compra)
        self.documento_confirmado(
            'COMPRA', [self.linea('NC_DESCUENTO_COMPRA', '1', '1', detalle_origen=detalle.id)],
            tipo=self.nota_credito, origen=Documento.objects.get(pk=compra), fecha='2026-01-15',
        )
        data = self.kardex()
        fila = data['movimientos'][-1]
        self.assertEqual(fila['naturaleza'], 'AJUSTE_VALOR')
        self.assertEqual(D(fila['saldo_valor']), D('2.00000000'))  # 3 − 1
        self.assertEqual(D(fila['cpp']), D('0.66666700'))
        self.assert_filas_contables(data)

    # 4
    def test_transferencia(self):
        TipoOperacion.objects.create(codigo_interno='TRANSFERENCIA_SALIDA', codigo_sunat='11', descripcion='TRANSF. SALIDA',
                                     naturaleza='SALIDA', origen_requerido='TRANSFERENCIA', es_interna=True)
        TipoOperacion.objects.create(codigo_interno='TRANSFERENCIA_ENTRADA', codigo_sunat='21', descripcion='TRANSF. ENTRADA',
                                     naturaleza='ENTRADA', origen_requerido='TRANSFERENCIA', es_interna=True)
        destino = Almacen.objects.create(empresa=self.empresa, establecimiento=self.establecimiento, codigo='ALM2', nombre='DESTINO')
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        transferencia = Transferencia.objects.create(empresa=self.empresa, almacen_origen=self.almacen,
                                                     almacen_destino=destino, fecha=date(2026, 1, 12), creado_por=self.usuario)
        TransferenciaDetalle.objects.create(transferencia=transferencia, producto=self.producto, cantidad_base=D('1'))
        confirmar_transferencia(transferencia_id=transferencia.id, empresa_id=self.empresa.id, usuario=self.usuario)

        consolidado = self.kardex()
        self.assertEqual(consolidado['resumen']['ambito_saldo'], 'EMPRESA')
        for fila in consolidado['movimientos'][-2:]:
            self.assertEqual((D(fila['saldo_cantidad']), D(fila['saldo_valor'])), (D('3'), D('2')))
        self.assert_filas_contables(consolidado)

        en_destino = self.kardex(almacen_id=destino.id)
        self.assertEqual(en_destino['resumen']['ambito_saldo'], 'ALMACEN')
        self.assertEqual(len(en_destino['movimientos']), 1)
        self.assertEqual(D(en_destino['movimientos'][0]['saldo_cantidad']), D('1'))
        self.assertEqual(D(en_destino['movimientos'][0]['saldo_valor']), D('0.66666700'))  # 1 × CPP global
        en_origen = self.kardex(almacen_id=self.almacen.id)
        self.assertEqual(D(en_origen['resumen']['saldo_final_cantidad']), D('2'))

    # 5
    def test_residuo(self):
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        self.venta('3')  # 3 × 0.666667 = 2.000001 → residuo −0.000001 absorbido
        data = self.kardex()
        self.assertTrue(any(f['generado_por_motor'] for f in data['movimientos']))
        self.assertEqual((D(data['resumen']['saldo_final_cantidad']), D(data['resumen']['saldo_final_valor'])),
                         (D('0'), D('0')))
        self.assert_filas_contables(data)

    # 6
    def test_extorno(self):
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        venta = self.venta('2')
        self.anular(venta)
        data = self.kardex()
        # El extorno (ENTRADA, prioridad 2) queda antes de la venta (prioridad 4) del mismo día.
        extorno = next(f for f in data['movimientos'] if f['es_extorno'])
        self.assertEqual(D(extorno['entrada_valor']), D('1.33333400'))  # 2 × 0.666667
        self.assertEqual(D(extorno['saldo_cantidad']), D('5'))
        self.assertEqual(D(extorno['saldo_valor']), D('3.33333400'))
        self.assertEqual((D(data['resumen']['saldo_final_cantidad']), D(data['resumen']['saldo_final_valor'])),
                         (D('3'), D('2.00000000')))
        self.assert_filas_contables(data)

    # 7
    def test_filtro_desde_fecha_intermedia(self):
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        self.venta('1', fecha='2026-01-20')
        data = self.kardex(fecha_desde='2026-01-15')
        self.assertEqual(len(data['movimientos']), 1)
        self.assertEqual(D(data['resumen']['saldo_inicial_cantidad']), D('3'))
        self.assertEqual(D(data['resumen']['saldo_inicial_valor']), D('2.00000000'))
        self.assertEqual(data['saldos_iniciales'][0]['cpp'], '0.66666700')
        fila = data['movimientos'][0]
        self.assertEqual(D(fila['saldo_cantidad']), D('2'))  # no arranca en 0
        self.assertEqual(D(fila['saldo_valor']), D('1.33333300'))  # 2.00000000 − 0.66666700
        self.assert_filas_contables(data)

    # 8
    def test_producto_con_movimientos_previos_sin_movimientos_en_rango(self):
        self.compra_confirmada('10', '10')
        self.compra_confirmada('10', '20', fecha='2026-01-25')  # posterior al rango
        data = self.kardex(fecha_desde='2026-01-15', fecha_hasta='2026-01-20')
        self.assertEqual(data['movimientos'], [])
        self.assertEqual(D(data['resumen']['saldo_inicial_valor']), D('100'))
        # Antes devolvía el saldo ACTUAL (20 u. / 300), ajeno al rango.
        self.assertEqual(D(data['resumen']['saldo_final_cantidad']), D('10'))
        self.assertEqual(D(data['resumen']['saldo_final_valor']), D('100'))
        self.assertEqual(D(data['resumen']['cpp_final']), D('10'))

    def test_filtro_por_tipo_no_oculta_movimientos_del_saldo(self):
        self.compra_confirmada('10', '10')
        self.venta('4')
        data = self.kardex(tipo_operacion_id=self.op['VENTA'].id)
        self.assertEqual(len(data['movimientos']), 1)
        self.assertEqual(D(data['movimientos'][0]['saldo_cantidad']), D('6'))  # antes: −4
        self.assertEqual(D(data['movimientos'][0]['saldo_valor']), D('60'))

    # 9
    def test_api_igual_a_excel(self):
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        self.venta('1', fecha='2026-01-20')
        for filtros in ({}, {'fecha_desde': '2026-01-11'}, {'periodo_id': self.enero.id}):
            with self.subTest(filtros=filtros):
                data, ws = self.kardex(**filtros), self.excel(**filtros)
                fila_excel = 8
                if filtros:
                    self.assertEqual(ws.cell(row=8, column=1).value, 'SALDO INICIAL')
                    self.assertEqual(D(str(ws.cell(row=8, column=12).value)), D(data['resumen']['saldo_inicial_cantidad']))
                    self.assertEqual(D(str(ws.cell(row=8, column=14).value)), D(data['resumen']['saldo_inicial_valor']))
                    fila_excel = 9
                for i, fila in enumerate(data['movimientos']):
                    for col, campo in ((12, 'saldo_cantidad'), (13, 'cpp'), (14, 'saldo_valor')):
                        self.assertAlmostEqual(ws.cell(row=fila_excel + i, column=col).value, float(fila[campo]), places=8)
                totales = fila_excel + len(data['movimientos'])
                self.assertEqual(ws.cell(row=totales, column=1).value, 'TOTALES DEL REPORTE')
                self.assertAlmostEqual(ws.cell(row=totales, column=14).value, float(data['resumen']['saldo_final_valor']), places=8)

    # 10
    def test_saldo_final_igual_a_fuente_contable(self):
        self.compra_confirmada('1', '1')
        self.compra_confirmada('2', '0.5', fecha='2026-01-11')
        self.venta('1', fecha='2026-01-20')
        resumen = self.kardex()['resumen']
        saldo, val = self.saldo(), self.ultima_valorizacion()
        self.assertEqual(D(resumen['saldo_final_cantidad']), saldo.cantidad_total)
        self.assertEqual(D(resumen['saldo_final_valor']), saldo.valor_total_pen)
        self.assertEqual(D(resumen['saldo_final_valor']), val.saldo_valor_empresa)
        self.assertEqual(D(resumen['cpp_final']), saldo.cpp_pen)

    def test_varios_productos_suma_saldos_por_producto(self):
        self.compra_confirmada('10', '10')
        self.documento_confirmado('COMPRA', [self.linea('COMPRA', '5', '3', producto=self.producto_2)])
        resumen = self.client.get(reverse('kardex-report'), {}, **self.headers).data['resumen']
        self.assertEqual(D(resumen['saldo_final_valor']), D('115'))
        self.assertEqual(D(resumen['saldo_final_cantidad']), D('15'))
        self.assertIsNone(resumen['cpp_final'])  # no existe un CPP común a dos productos
