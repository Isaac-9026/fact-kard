"""
F131-001 — el XLSX del Formato 13.1 no traía la fila de saldo final del período.

La pantalla termina cada establecimiento con `reporte['saldo_final']` (cierre vigente o, sin
cierre, cantidad × CPP global a la fecha fin). El XLSX terminaba con la última fila de
movimientos y pasaba directo al pie "FUENTE DEL SALDO FINAL". En un anexo esa última fila
conserva el CPP del momento de su último movimiento; si después otro establecimiento mueve el
CPP global, el saldo final oficial es otro (caso auditado: 11 × 10.27263155 = 112.998947 frente
al oficial 11 × 10.221167 = 112.432837).

Escenario: saldo inicial y compra en el principal, transferencia de 11 u. al anexo y luego una
compra más cara en el principal (el CPP global cambia después del último movimiento del anexo).
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.urls import reverse

from apps.empresas.models import CierrePeriodo
from apps.inventario.models import SaldoInicial, Transferencia, TransferenciaDetalle
from apps.inventario.services import confirmar_saldo_inicial, confirmar_transferencia
from apps.reportes import tests_b4_conciliacion_integral as b4
from apps.reportes.formato131_excel import ETIQUETA_SALDO_FINAL

D = Decimal


class SaldoFinalExcelTests(b4.ConciliacionIntegralTests):

    def test_conciliacion_integral(self):  # el escenario heredado ya se prueba en su módulo
        pass

    def escenario(self):
        saldo = SaldoInicial.objects.create(empresa=self.empresa, periodo=self.enero, almacen=self.almacen,
                                            producto=self.producto, cantidad=D('50'), valor_pen=D('500.00'),
                                            cpp_inicial=D('10'), creado_por=self.usuario)
        confirmar_saldo_inicial(saldo_id=saldo.id, empresa_id=self.empresa.id, usuario=self.usuario)
        self.compra_confirmada('50', '10.40', fecha='2026-01-05')           # 100 u. → CPP 10.20
        t = Transferencia.objects.create(empresa=self.empresa, almacen_origen=self.almacen, almacen_destino=self.almacen_b,
                                         fecha=date(2026, 1, 10), creado_por=self.usuario)
        TransferenciaDetalle.objects.create(transferencia=t, producto=self.producto, cantidad_base=D('11'))
        confirmar_transferencia(transferencia_id=t.id, empresa_id=self.empresa.id, usuario=self.usuario)
        self.compra_confirmada('10', '13.30', fecha='2026-01-20')           # el CPP global cambia después

    def exportar(self, **params):
        params.setdefault('periodo_id', self.enero.id)
        resp = self.client.get(reverse('formato-131-excel'), params, **self.headers)
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    def pantalla(self, **params):
        params.setdefault('periodo_id', self.enero.id)
        resp = self.client.get(reverse('formato-131'), params, **self.headers)
        self.assertEqual(resp.status_code, 200)
        return {r['cabecera']['establecimiento']['codigo_sunat']: r for r in resp.data['reportes']}

    def hoja_de(self, libro, codigo):
        return next(libro[n] for n in libro.sheetnames if str(libro[n]['B5'].value).startswith(f'{codigo} - '))

    @staticmethod
    def filas(hoja):
        return [[hoja.cell(r, c).value for c in range(1, 15)] for r in range(1, hoja.max_row + 1)]

    def bloque(self, fila):
        return [D(str(v)) for v in fila[11:14]]

    def verificar_hoja(self, hoja, reporte):
        filas = self.filas(hoja)
        etiquetas = [f[0] for f in filas]
        self.assertEqual(etiquetas.count(ETIQUETA_SALDO_FINAL), 1)  # una sola fila de saldo final
        i_final = etiquetas.index(ETIQUETA_SALDO_FINAL)
        i_ultimo_mov = 13 + len(reporte['movimientos']) - 1  # índice 0-based de la última fila de movimiento
        self.assertEqual(i_final, i_ultimo_mov + 1)          # justo después de los movimientos
        self.assertEqual(etiquetas[i_final + 2], 'FUENTE DEL SALDO FINAL')
        ultimo = reporte['movimientos'][-1]['saldo_final']
        self.assertEqual(self.bloque(filas[i_ultimo_mov]), [D(ultimo[k]) for k in ('cantidad', 'costo_unitario', 'costo_total')])
        final = reporte['saldo_final']
        self.assertEqual(self.bloque(filas[i_final]), [D(final[k]) for k in ('cantidad', 'costo_unitario', 'costo_total')])
        return self.bloque(filas[i_ultimo_mov]), self.bloque(filas[i_final])

    def test_anexo_termina_con_el_saldo_final_oficial_y_no_con_el_ultimo_movimiento(self):
        self.escenario()
        self.client.post(reverse('periodo-cerrar', args=[self.enero.id]), format='json', **self.headers)
        cierre = CierrePeriodo.objects.get(periodo=self.enero, vigente=True)
        reportes = self.pantalla()
        libro = self.exportar()
        intermedio, final = self.verificar_hoja(self.hoja_de(libro, '0001'), reportes['0001'])
        # Último movimiento del anexo: 11 × 10.20 (CPP al transferir). Oficial: 11 × CPP final.
        self.assertEqual(intermedio, [D('11'), D('10.2'), D('112.2')])
        detalle = cierre.detalles.get(producto=self.producto, establecimiento=self.sucursal)
        self.assertEqual(final[0], D('11'))
        self.assertEqual(final[2], D(detalle.saldo_valor_pen))
        self.assertEqual(final[1], D(reportes['0001']['saldo_final']['costo_unitario']))
        self.assertNotEqual(final, intermedio)  # antes el XLSX terminaba con 11 × 10.20 = 112.20
        # Principal también termina con su saldo oficial.
        self.verificar_hoja(self.hoja_de(libro, '0000'), reportes['0000'])
        # El total de los establecimientos concilia con el cierre (cantidad y valor).
        totales = [self.bloque(self.filas(h)[[f[0] for f in self.filas(h)].index(ETIQUETA_SALDO_FINAL)]) for h in libro.worksheets]
        self.assertEqual(sum(t[0] for t in totales), sum(D(x.saldo_cantidad) for x in cierre.detalles.all()))
        self.assertEqual(sum(t[2] for t in totales), sum(D(x.saldo_valor_pen) for x in cierre.detalles.all()))
        self.assertEqual(sum(t[0] for t in totales), D('110'))

    def test_periodo_abierto_tambien_cierra_con_el_saldo_final_calculado(self):
        self.escenario()
        reportes = self.pantalla()
        libro = self.exportar()
        for codigo in ('0000', '0001'):
            self.verificar_hoja(self.hoja_de(libro, codigo), reportes[codigo])
        hoja = self.hoja_de(libro, '0001')
        fuente = [f for f in self.filas(hoja) if f[0] == 'FUENTE DEL SALDO FINAL'][0][1]
        self.assertTrue(fuente.startswith('CALCULADO'))

    def test_filtros_por_producto_y_establecimiento(self):
        self.escenario()
        self.client.post(reverse('periodo-cerrar', args=[self.enero.id]), format='json', **self.headers)
        libro = self.exportar(establecimiento_id=self.sucursal.id, producto_id=self.producto.id)
        self.assertEqual(len(libro.worksheets), 1)
        reportes = self.pantalla(establecimiento_id=self.sucursal.id, producto_id=self.producto.id)
        self.verificar_hoja(libro.worksheets[0], reportes['0001'])
        libro = self.exportar(producto_id=self.producto.id)
        self.assertEqual(len(libro.worksheets), 2)
