"""
QA-003 — Formato 13.1: saldo intermedio del establecimiento destino de una transferencia.

Causa: la entrada de la transferencia (prioridad 2 del día) se registra antes que su salida
(prioridad 4). Entre ambas las cantidades por establecimiento no suman las de la empresa y
el reparto del Bloque 5 no se calcula para ese instante; la fila del destino caía en
cantidad × CPP de la empresa EN ESE INSTANTE. Si el mismo día hay un ajuste de valor
(NC de descuento, prioridad 3) el CPP aún no lo refleja: entrada 10 × 12.25 = 122.50 pero
saldo 10 × 12.50 = 125.00, y al cierre volvía a 12.25.

Corrección: en ese instante el saldo del destino es su valor anterior + lo que entra (al costo
de la transferencia). No cambia el CPP global, el Kardex ni el saldo final.
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.urls import reverse

from apps.documentos.models import Documento, DocumentoDetalle
from apps.inventario.models import Almacen, Transferencia, TransferenciaDetalle
from apps.inventario.services import confirmar_transferencia
from apps.reportes import tests_b4_conciliacion_integral as b4
from apps.reportes.formato131_service import generar_formato_131

D = Decimal


class TransferenciaSaldoIntermedioTests(b4.ConciliacionIntegralTests):

    def test_conciliacion_integral(self):  # el escenario heredado ya se prueba en su módulo
        pass

    def transferir(self, dia, cantidad, destino=None):
        t = Transferencia.objects.create(empresa=self.empresa, almacen_origen=self.almacen, almacen_destino=destino or self.almacen_b,
                                         fecha=date(2026, 1, dia), creado_por=self.usuario)
        TransferenciaDetalle.objects.create(transferencia=t, producto=self.producto, cantidad_base=D(cantidad))
        confirmar_transferencia(transferencia_id=t.id, empresa_id=self.empresa.id, usuario=self.usuario)

    def escenario(self, con_ajuste=True):
        compra = self.compra_confirmada('10', '12.5', fecha='2026-01-05')
        if con_ajuste:
            self.documento_confirmado(
                'COMPRA', [self.linea('NC_DESCUENTO_COMPRA', '1', '2.50',
                                      detalle_origen=DocumentoDetalle.objects.get(documento_id=compra).id)],
                tipo=self.nota_credito, origen=Documento.objects.get(pk=compra), fecha='2026-01-07',
            )
        self.transferir(7, '10')

    def reportes(self):
        return {r['cabecera']['establecimiento']['id']: r
                for r in generar_formato_131(empresa_id=self.empresa.id, periodo_id=self.enero.id)['reportes']}

    def kardex(self):
        return self.client.get(reverse('kardex-report'), {'producto_id': self.producto.id}, **self.headers).data

    def test_entrada_destino_y_saldo_intermedio_con_el_cpp_de_la_transferencia(self):
        self.escenario()
        destino = self.reportes()[self.sucursal.id]
        fila = destino['movimientos'][0]
        self.assertEqual(fila['tipo_operacion_interno'], 'TRANSFERENCIA_ENTRADA')
        self.assertEqual({k: D(v) for k, v in fila['entrada'].items()},
                         {'cantidad': D('10'), 'costo_unitario': D('12.25'), 'costo_total': D('122.50')})
        # Antes: 10 / 12.50 / 125.00.
        self.assertEqual({k: D(v) for k, v in fila['saldo_final'].items()},
                         {'cantidad': D('10'), 'costo_unitario': D('12.25'), 'costo_total': D('122.50')})
        # Saldo anterior (0) + entrada = saldo de la fila; y coincide con el saldo final del destino.
        self.assertEqual(D(destino['saldo_inicial_valorizado']['costo_total']) + D(fila['entrada']['costo_total']),
                         D(fila['saldo_final']['costo_total']))
        self.assertEqual(D(destino['saldo_final']['costo_total']), D('122.50'))
        self.assertEqual(D(destino['saldo_final']['costo_unitario']), D('12.25'))

    def test_mismo_cpp_global_kardex_y_saldo_final_sin_cambios(self):
        self.escenario()
        kardex = self.kardex()
        ultimo = kardex['movimientos'][-1]
        self.assertEqual(D(ultimo['saldo_cantidad']), D('10'))
        self.assertEqual(D(ultimo['saldo_valor']), D('122.50'))
        self.assertEqual(D(ultimo['cpp']), D('12.25'))
        reportes = self.reportes()
        origen, destino = reportes[self.establecimiento.id], reportes[self.sucursal.id]
        self.assertEqual(D(origen['saldo_final']['cantidad']) + D(destino['saldo_final']['cantidad']), D('10'))
        self.assertEqual(D(origen['saldo_final']['costo_total']) + D(destino['saldo_final']['costo_total']), D('122.50'))
        salida = origen['movimientos'][-1]
        self.assertEqual(salida['tipo_operacion_interno'], 'TRANSFERENCIA_SALIDA')
        self.assertEqual(D(salida['salida']['costo_unitario']), D('12.25'))
        self.assertEqual(D(salida['saldo_final']['costo_total']), D('0'))

    def test_excel_13_1_consistente_con_el_reporte(self):
        self.escenario()
        resp = self.client.get(reverse('formato-131-excel'), {'periodo_id': self.enero.id}, **self.headers)
        self.assertEqual(resp.status_code, 200)
        libro = openpyxl.load_workbook(io.BytesIO(resp.content))
        hoja = next(libro[n] for n in libro.sheetnames if '0001 - SUCURSAL' in str(libro[n]['B5'].value))
        fila = [hoja.cell(14, c).value for c in range(5, 15)]
        self.assertEqual(fila[0], '21')  # Tabla 12: transferencia entrada
        entrada, saldo = fila[1:4], fila[7:10]
        self.assertEqual([D(str(v)) for v in entrada], [D('10'), D('12.25'), D('122.50')])
        self.assertEqual([D(str(v)) for v in saldo], [D('10'), D('12.25'), D('122.50')])

    def test_sin_ajuste_el_mismo_dia_no_cambia_nada(self):
        self.escenario(con_ajuste=False)
        fila = self.reportes()[self.sucursal.id]['movimientos'][0]
        self.assertEqual([D(fila['saldo_final'][k]) for k in ('cantidad', 'costo_unitario', 'costo_total')],
                         [D('10'), D('12.5'), D('125')])

    def test_compra_del_mismo_dia_registrada_despues_de_la_transferencia(self):
        # Otra forma de llegar al mismo síntoma: 10 u. a 12.50 en ALM1; el día 7 se transfiere a
        # ALM2 y luego se registra una compra del mismo día de 10 u. a 12.00. La compra (prioridad 2,
        # secuencia posterior) queda entre la entrada y la salida de la transferencia: CPP 12.25.
        self.compra_confirmada('10', '12.5', fecha='2026-01-05')
        self.transferir(7, '10')
        self.compra_confirmada('10', '12', fecha='2026-01-07')
        reportes = self.reportes()
        fila = reportes[self.sucursal.id]['movimientos'][0]
        self.assertEqual(fila['tipo_operacion_interno'], 'TRANSFERENCIA_ENTRADA')
        entrada = [D(fila['entrada'][k]) for k in ('cantidad', 'costo_unitario', 'costo_total')]
        saldo = [D(fila['saldo_final'][k]) for k in ('cantidad', 'costo_unitario', 'costo_total')]
        self.assertEqual(saldo[0], D('10'))
        self.assertEqual(saldo[2], entrada[2], f'entrada {entrada} / saldo {saldo}')
        self.assertEqual(saldo[1], entrada[1])
        kardex = self.kardex()['movimientos'][-1]
        total = sum(D(r['saldo_final']['costo_total']) for r in reportes.values())
        self.assertEqual(total, D(kardex['saldo_valor']))

    def test_transferencia_dentro_del_mismo_establecimiento_cuadra_fila_a_fila(self):
        # Ambos almacenes en el mismo establecimiento: cada fila = saldo anterior ± movimiento.
        alm2 = Almacen.objects.create(empresa=self.empresa, establecimiento=self.establecimiento, codigo='ALM2', nombre='ALM2')
        compra = self.compra_confirmada('10', '12.5', fecha='2026-01-05')
        self.documento_confirmado(
            'COMPRA', [self.linea('NC_DESCUENTO_COMPRA', '1', '2.50',
                                  detalle_origen=DocumentoDetalle.objects.get(documento_id=compra).id)],
            tipo=self.nota_credito, origen=Documento.objects.get(pk=compra), fecha='2026-01-07',
        )
        self.transferir(7, '10', destino=alm2)
        reporte = self.reportes()[self.establecimiento.id]
        anterior = D(reporte['saldo_inicial_valorizado']['costo_total'])
        for fila in reporte['movimientos']:
            entra = D(fila['entrada']['costo_total']) if fila['entrada'] else D(0)
            sale = D(fila['salida']['costo_total']) if fila['salida'] else D(0)
            self.assertEqual(anterior + entra - sale, D(fila['saldo_final']['costo_total']), fila['tipo_operacion_interno'])
            anterior = D(fila['saldo_final']['costo_total'])
        self.assertEqual(anterior, D('122.50'))
        self.assertEqual(D(reporte['saldo_final']['costo_total']), D('122.50'))
